"""
Planilha "foto enviada × avatar gerado × está parecido?" (RF40).

Uso: python3 build-likeness-xlsx.py <faces-sel.json> <fotos> <renders-antes> <renders-depois> <likeness.json> <julgamento.json> <saida.xlsx>
  renders-antes/depois: saídas de validate.mjs antes e depois do ajuste do controle de qualidade;
  likeness.json: saída de eval-likeness.py (renders-depois);
  julgamento.json: avaliação visual por pessoa (cabelo, pele, cabeça, parecido, observação).
"""
import io
import json
import sys

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

sel_p, photos, before, after, like_p, judge_p, out = sys.argv[1:8]
sel = json.load(open(sel_p, encoding="utf-8"))
res_before = {r["name"]: r for r in json.load(open(f"{before}/results.json"))}
res_after = {r["name"]: r for r in json.load(open(f"{after}/results.json"))}
like = json.load(open(like_p, encoding="utf-8")); L = {r["nome"]: r for r in like["linhas"]}
J = json.load(open(judge_p, encoding="utf-8"))

H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F2A44"); WRAP = Alignment(wrap_text=True, vertical="top")
OK = PatternFill("solid", fgColor="DCF2E3"); BAD = PatternFill("solid", fgColor="F9DEDC"); MID = PatternFill("solid", fgColor="FFF1CC")


def img_cell(ws, path, anchor, size=150):
    im = Image.open(path).convert("RGB"); im.thumbnail((size, size))
    buf = io.BytesIO(); im.save(buf, "JPEG", quality=82); buf.seek(0)
    ws.add_image(XlImage(buf), anchor)


def status(r):
    if r["state"]["status"] == "built":
        return "Gerado"
    blocks = [i["code"] for p in r["state"].get("photos", []) for i in p["issues"] if i["severity"] == "block"]
    return "Recusado: " + ", ".join(blocks)


wb = Workbook()
ws = wb.active; ws.title = "Semelhança"
cols = ["Pessoa", "Foto enviada (upload)", "Avatar gerado — frente", "Avatar gerado — 3/4", "Está parecido? (sim/não)", "Observação",
        "Antes do ajuste", "Depois do ajuste", "Avisos da foto", "Cabelo certo?", "Tom de pele certo?", "Formato da cabeça",
        "Reconhecedor SFace, frente (mesma pessoa ≥ 0,363)", "Reconhecedor ArcFace, frente (referência ≥ 0,30)",
        "O avatar é mais parecido com a própria foto do que com as outras 15?", "Autor da foto", "Página original (Flickr)", "Licença"]
ws.append(cols)
widths = [11, 22, 22, 22, 14, 48, 26, 12, 34, 16, 14, 18, 16, 16, 20, 18, 34, 10]
for i, w in enumerate(widths, 1):
    c = ws.cell(row=1, column=i); c.font = H; c.fill = HF; c.alignment = WRAP; ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "B2"
sim = 0
for k, s in enumerate(sel, start=2):
    n = s["nome"]; ra = res_after[n]; rb = res_before[n]; j = J[n]; lv = (L.get(n) or {}).get("vistas", {}).get("front", {})
    warns = [i["code"] for p in ra["state"].get("photos", []) for i in p["issues"] if i["severity"] == "warn"]
    ident = "" if not lv else ("Sim (1º de 16)" if lv.get("posicaoEntreFotos") == 1 else f"Não ({lv.get('posicaoEntreFotos')}º de 16)")
    ws.append([n, "", "", "", j["parecido"], j["obs"], status(rb), status(ra), ", ".join(warns) or "—", j["cabelo"], j["pele"], j["cabeca"],
               lv.get("sface", "—"), lv.get("arcface", "—"), ident or "—", s["Author"], s["OriginalLandingURL"], "CC BY 2.0"])
    sim += j["parecido"] == "Sim"
    img_cell(ws, f"{photos}/{n}.jpg", f"B{k}")
    if ra["shots"].get("front"):
        img_cell(ws, f"{after}/{ra['shots']['front']}", f"C{k}"); img_cell(ws, f"{after}/{ra['shots']['left34']}", f"D{k}")
    ws.row_dimensions[k].height = 118
    for c in range(1, len(cols) + 1):
        ws.cell(row=k, column=c).alignment = WRAP
    ws.cell(row=k, column=5).fill = OK if j["parecido"] == "Sim" else BAD
    ws.cell(row=k, column=17).hyperlink = s["OriginalLandingURL"]
    for col, v in ((10, j["cabelo"]), (11, j["pele"])):
        ws.cell(row=k, column=col).fill = OK if v.startswith("Sim") else (MID if v.startswith("Parcial") else BAD)

r = wb.create_sheet("Resumo", 0)
built_b = sum(1 for x in res_before.values() if x["state"]["status"] == "built")
built_a = sum(1 for x in res_after.values() if x["state"]["status"] == "built")
fronts = [v["vistas"]["front"] for v in like["linhas"] if v["vistas"].get("front", {}).get("rostoDetectado")]
for a, b in [
    ("Semelhança do Avatar 3D com a foto enviada (27/09/2026)", ""),
    ("Fotos testadas (adultos, uma foto só por pessoa)", len(sel)),
    ("Avatares gerados antes do ajuste do controle de qualidade", f"{built_b} de {len(sel)}"),
    ("Avatares gerados depois do ajuste", f"{built_a} de {len(sel)}"),
    ("Está parecido? (julgamento visual: alguém que conhece a pessoa a reconheceria de imediato)", f"{sim} de {len(sel)}"),
    ("Reconhecedor facial diz que o rosto do avatar é da mesma pessoa (SFace, frente)", f"{sum(1 for f in fronts if f['sface'] >= 0.363)} de {len(fronts)}"),
    ("Avatar mais parecido com a própria foto do que com as outras 15", f"{sum(1 for f in fronts if f['posicaoEntreFotos'] == 1)} de {len(fronts)}"),
    ("Semelhança média entre pessoas diferentes deste teste (SFace)", like["impostores"]["sfaceMedia"]),
    ("", ""),
    ("Leitura", "O miolo do rosto é a própria foto aplicada como textura, por isso o reconhecedor facial confirma a identidade. "
                "O resto não acompanha: o cabelo vira uma calota genérica (cor às vezes certa, forma quase nunca), turbante, lenço e "
                "barba longa somem, o crânio e as orelhas são os do manequim e a textura tem manchas onde a foto tinha sombra. Por isso "
                "nenhum avatar foi julgado 'extremamente parecido'. O ajuste de hoje só evita recusar fotos boas; a semelhança "
                "depende de trocar a reconstrução (ver docs/avatar3d/proposta-profissional-2026-09-27.md)."),
]:
    r.append([a, b])
r.column_dimensions["A"].width = 80; r.column_dimensions["B"].width = 100; r["A1"].font = Font(bold=True, size=14)
r.cell(row=10, column=2).alignment = WRAP

m = wb.create_sheet("Método")
for line in [
    "Fotos: Open Images V7 (Google), partições de validação e teste, fotos do Flickr sob CC BY 2.0. Filtro: exatamente um rosto humano, "
    "não encoberto, não cortado, não desenho, ocupando de 22% a 60% da largura. Escolha manual de 16 adultos de frente, com tons de pele variados.",
    "Excluídas de propósito: crianças, pessoas públicas reconhecíveis e fotos de perfil ou muito viradas.",
    "Avatar: o pipeline real do app (MediaPipe no navegador + atlas + busto), pelo laboratório /lab/avatar e scripts/avatar3d/validate.mjs, uma foto por pessoa.",
    "Antes do ajuste: 8 fotos recusadas, quase todas por 'rosto encoberto' (barba, sombra e variação de tom contam como obstrução). "
    "Depois: esse aviso não recusa mais, e a luz lateral só recusa acima de 2,6 de desequilíbrio.",
    "Medida objetiva: scripts/avatar3d/eval-likeness.py com SFace (OpenCV Zoo) e ArcFace (ONNX Model Zoo), frente e 3/4.",
    "Julgamento visual: feito por quem montou o teste, olhando foto e avatar lado a lado; critério na aba Resumo.",
]:
    m.append([line])
m.column_dimensions["A"].width = 170
wb.save(out)
print("ok", out, "parecidos:", sim)
