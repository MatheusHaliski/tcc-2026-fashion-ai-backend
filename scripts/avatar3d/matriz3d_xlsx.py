#!/usr/bin/env python3
"""MATRIZ-3D (#237) · exporta as tabelas versionadas para planilha.

    MATRIZ3D_WRITE=1 npx vitest run lib/avatar3d/matriz3d     # 1. regrava os CSV/JSON a partir da implementação
    python3 scripts/avatar3d/matriz3d_xlsx.py                 # 2. gera docs/avatar3d/matriz-3d/matriz-3d.xlsx

Abas: Leia-me, Matriz principal (uma linha por subcategoria), Componentes (uma linha por subcategoria × componente),
Biblioteca (contratos, dependências e testes de cada componente), Achados e Legenda. A planilha é só leitura: a fonte de
verdade são os CSV gerados pelo teste (editar a planilha não muda a matriz).
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DIR = ROOT / "docs/avatar3d/matriz-3d"
OUT = DIR / "matriz-3d.xlsx"

ESTADO_COR = {
    "Implementado e validado": "C8E6C9", "Implementado, sem validação": "FFF3C4", "Parcial": "FFE0B2",
    "Ausente": "FFCDD2", "Não aplicável": "ECEFF1",
}
ICONE_COR = {"✔": ESTADO_COR["Implementado e validado"], "◐": ESTADO_COR["Implementado, sem validação"], "◑": ESTADO_COR["Parcial"], "✖": ESTADO_COR["Ausente"]}
HEAD = PatternFill("solid", fgColor="263238")


def read_csv(name: str) -> list[list[str]]:
    with (DIR / name).open(encoding="utf-8", newline="") as fh:
        return list(csv.reader(fh))


def sheet(wb: Workbook, title: str, rows: list[list[str]], widths: dict[int, int] | None = None, color_col: str | None = None, icon_cols: bool = False):
    ws = wb.create_sheet(title)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2" if not icon_cols else "E2"
    ws.auto_filter.ref = ws.dimensions
    header = rows[0]
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(60, max(10, len(h) + 4)))
    color_idx = header.index(color_col) + 1 if color_col and color_col in header else None
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if color_idx:
            fill = ESTADO_COR.get(str(row[color_idx - 1].value))
            if fill:
                row[color_idx - 1].fill = PatternFill("solid", fgColor=fill)
        if icon_cols:
            for c in row:
                v = str(c.value or "")
                first = v[:1]
                if first in ICONE_COR and "\n" not in v:
                    c.fill = PatternFill("solid", fgColor=ICONE_COR[first])
    return ws


def main() -> int:
    if not (DIR / "matriz-principal.csv").exists():
        print("rode antes: MATRIZ3D_WRITE=1 npx vitest run lib/avatar3d/matriz3d", file=sys.stderr)
        return 2
    meta = json.loads((DIR / "matriz-3d.json").read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Leia-me"
    v = meta["versao"]
    for line in [
        ["MATRIZ-3D — rastreabilidade da vestimenta 3D do FashionAI"],
        [f"Versões: taxonomia {v['taxonomia']} · molde {v['molde']} · contrato {v['contrato']} · corpo {v['corpo']}"],
        ["Fonte de verdade: os CSV de docs/avatar3d/matriz-3d, gerados pela execução real dos construtores (lib/avatar3d/matriz3d)."],
        ["Atualizar: MATRIZ3D_WRITE=1 npx vitest run lib/avatar3d/matriz3d && python3 scripts/avatar3d/matriz3d_xlsx.py"],
        ["Abas: Matriz principal (subcategoria × grupo) · Componentes (subcategoria × componente) · Biblioteca · Achados · Legenda"],
        [""],
        ["Estados"],
        *[[f"{k}"] for k in ESTADO_COR],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=14)
    for i, k in enumerate(ESTADO_COR, 8):
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=ESTADO_COR[k])
    ws.column_dimensions["A"].width = 120

    principal = read_csv("matriz-principal.csv")
    widths = {i: 18 for i in range(1, 16)}
    widths.update({i: 48 for i in range(16, 29)})
    sheet(wb, "Matriz principal", principal, widths, icon_cols=True)
    sheet(wb, "Componentes", read_csv("componentes.csv"), {1: 16, 2: 18, 3: 26, 4: 36, 5: 18, 6: 24, 7: 14, 8: 22, 9: 40, 10: 48, 11: 48, 12: 26, 13: 48, 14: 24, 15: 60}, color_col="estado")
    sheet(wb, "Biblioteca", read_csv("biblioteca.csv"), {1: 26, 2: 18, 3: 40, 4: 30, 5: 48, 6: 48, 7: 60, 8: 30, 9: 60, 10: 30, 11: 8, 12: 10, 13: 60, 14: 8})
    sheet(wb, "Achados", read_csv("achados.csv"), {1: 26, 2: 44, 3: 90, 4: 9})
    leg = wb.create_sheet("Legenda")
    for row in [["Ícone", "Estado"], ["✔", "Implementado e validado"], ["◐", "Implementado, sem validação"], ["◑", "Parcial"], ["✖", "Ausente"], ["—", "Não aplicável"]]:
        leg.append(row)
    for i, k in enumerate(ESTADO_COR, 2):
        leg.cell(row=i, column=2).fill = PatternFill("solid", fgColor=ESTADO_COR[k])
    leg.column_dimensions["B"].width = 32
    wb.save(OUT)
    print(f"{OUT.relative_to(ROOT)}: {len(principal) - 1} subcategorias")
    return 0


if __name__ == "__main__":
    sys.exit(main())
