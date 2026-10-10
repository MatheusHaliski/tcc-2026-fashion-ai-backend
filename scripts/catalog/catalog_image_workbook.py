"""Exportação XLSX do inventário de imagens; não consulta APIs, baixa fotos ou altera o banco."""
from __future__ import annotations

from collections import Counter
from datetime import datetime
import json
import os
from pathlib import Path
import re
from typing import Any, Iterable, Mapping
from urllib.parse import parse_qsl, urlsplit
from uuid import uuid4
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = (
    "Nome da peça", "Marca", "Imagem URL", "Pipeline aplicado (sim/não)",
    "Pipeline após a execução", "Imagem original URL", "Imagem processada / recorte",
    "Status de processamento", "Motivo / observações", "Origem dos dados", "Versão do pipeline",
    "ID do produto", "ID da imagem", "Recorte JSON", "Erro",
)
UNKNOWN = "Não verificado"
STATUS_COLORS = {"Sim": "E0F0EE", "Não": "FBE7EE", UNKNOWN: "F1F0EA"}
URL_PATTERN = re.compile(r"https?://[^\s<>\"']+", re.IGNORECASE)
SECRET_QUERY_KEYS = {"password", "passwd", "secret", "api_key", "apikey", "access_token", "refresh_token", "token", "authorization", "x-amz-credential", "x-amz-signature"}
OMITTED_URL = "[URL omitida: credenciais]"
SNAPSHOT_NOTE = (
    "A coluna 'Pipeline aplicado (sim/não)' descreve o estado anterior à execução registrado pela fonte. "
    "Um snapshot sem metadados do pipeline recebe 'Não verificado'; ele não comprova o estado atual do banco. "
    "A coluna 'Pipeline após a execução' é separada para não substituir essa evidência histórica."
)
CROP_NOTE = (
    "SEMANTIC_CROP aplica o recorte na exibição usando a URL original e os metadados JSON; a URL pode permanecer "
    "igual, sem existir uma segunda imagem no storage. PROCESSED usa uma saída persistida. "
    "A existência de uma URL ou de um nome de versão, isoladamente, não comprova processamento nem revisão visual."
)


def _status(value: bool | None) -> str:
    if value is True:
        return "Sim"
    if value is False:
        return "Não"
    if value is None:
        return UNKNOWN
    raise ValueError("Os estados standardized_before/standardized_after devem ser bool ou None.")


def _has_credentials(value: str) -> bool:
    try:
        parsed = urlsplit(value)
        return parsed.username is not None or parsed.password is not None or any(
            key.lower() in SECRET_QUERY_KEYS for key, _ in parse_qsl(parsed.query)
        )
    except ValueError:
        return True


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    return URL_PATTERN.sub(lambda match: OMITTED_URL if _has_credentials(match.group()) else match.group(), value)[:32767]


def _url(value: Any) -> tuple[str, str | None]:
    raw = str(value or "").strip()
    if not raw:
        return "", None
    if _has_credentials(raw):
        return OMITTED_URL, None
    try:
        parsed = urlsplit(raw)
        valid = parsed.scheme.lower() in {"https", "http"} and bool(parsed.hostname)
    except ValueError:
        valid = False
    return _text(raw), raw if valid else None


def _cell(sheet, row: int, column: int, value: Any, *, link: str | None = None):
    cell = sheet.cell(row, column, _text(value))
    # Strings no XLSX, inclusive =/+/-/@, nunca fórmulas vindas do inventário.
    cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if link:
        cell.hyperlink = link
        cell.font = Font(color="145E83", underline="single")
    return cell


def _processed_link(record: Mapping[str, Any], workbook_path: Path) -> tuple[str, str | None]:
    output = record.get("output_path")
    if output:
        output_text = str(output)
        if "://" in output_text:
            return _url(output_text)
        # Saídas locais do lote são links relativos, para poder acompanhar o XLSX no pacote de entrega.
        local = Path(output_text)
        local = local if local.is_absolute() else local.resolve()
        if local.is_file():
            relative = Path(os.path.relpath(local, workbook_path.parent.resolve())).as_posix()
            return relative, relative
        return _text(output_text), None
    current = record.get("current_url") or record.get("image_url")
    original = record.get("source_url")
    return _url(current) if current and current != original else ("", None)


def _generated_at(summary: Mapping[str, Any]) -> datetime:
    value = summary.get("generated_at")
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            value = None
    zone = ZoneInfo("America/Sao_Paulo")
    if not isinstance(value, datetime):
        return datetime.now(zone)
    return value.replace(tzinfo=zone) if value.tzinfo is None else value.astimezone(zone)


def write_workbook(records: Iterable[Mapping[str, Any]], path: str | Path,
                   summary: Mapping[str, Any] | None = None) -> Path:
    """Gera o relatório e devolve seu Path; os estados informados não são inferidos a partir das URLs."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    summary = summary or {}
    book = Workbook()
    sheet = book.active
    sheet.title = "Acervo"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"
    sheet.row_dimensions[1].height = 36
    widths = (46, 24, 54, 28, 28, 54, 54, 28, 72, 28, 32, 40, 40, 70, 60)
    for column, (header, width) in enumerate(zip(HEADERS, widths), 1):
        cell = _cell(sheet, 1, column, header)
        cell.fill = PatternFill("solid", fgColor="1F504D")
        cell.font = Font(color="FFFFFF", bold=True)
        sheet.column_dimensions[cell.column_letter].width = width

    before_counts: Counter[str] = Counter()
    after_counts: Counter[str] = Counter()
    scope_counts: Counter[str] = Counter()
    processing_counts: Counter[str] = Counter()
    product_ids: set[str] = set()
    image_ids: set[str] = set()
    error_count = 0
    count = 0
    for count, record in enumerate(records, 1):
        if count >= 1048576:
            raise ValueError("O inventário ultrapassa o limite de linhas de uma aba XLSX.")
        before = _status(record.get("standardized_before"))
        after = _status(record.get("standardized_after"))
        current_text, current_link = _url(record.get("current_url") or record.get("image_url") or record.get("source_url"))
        source_text, source_link = _url(record.get("source_url"))
        processed_text, processed_link = _processed_link(record, path)
        scope = record.get("source_scope") or record.get("origin") or "Não informado"
        values = (
            record.get("product_name"), record.get("brand"), current_text, before, after,
            source_text, processed_text, record.get("processing_status"), record.get("status_note"),
            scope, record.get("pipeline_version"), record.get("product_id"), record.get("image_id"),
            record.get("crop_json", record.get("crop")), record.get("error"),
        )
        row = count + 1
        for column, value in enumerate(values, 1):
            link = {3: current_link, 6: source_link, 7: processed_link}.get(column)
            cell = _cell(sheet, row, column, value, link=link)
            if column in (4, 5):
                cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[value])
                cell.font = Font(bold=True, color="1A3632")
        sheet.row_dimensions[row].height = 48
        before_counts[before] += 1
        after_counts[after] += 1
        scope_counts[_text(scope)] += 1
        processing_counts[_text(record.get("processing_status")) or "Não informado"] += 1
        if record.get("product_id"):
            product_ids.add(_text(record["product_id"]))
        if record.get("image_id"):
            image_ids.add(_text(record["image_id"]))
        error_count += bool(record.get("error"))

    if count:
        table = Table(displayName="InventarioImagens", ref=f"A1:O{count + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
                                             showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    sheet.auto_filter.ref = f"A1:O{count + 1}"

    overview = book.create_sheet("Resumo", 0)
    overview.sheet_view.showGridLines = False
    overview.freeze_panes = "A2"
    overview.column_dimensions["A"].width = 42
    overview.column_dimensions["B"].width = 110
    overview.append(["Inventário de imagens do catálogo", "Método e evidências"])
    for cell in overview[1]:
        cell.fill = PatternFill("solid", fgColor="1F504D")
        cell.font = Font(color="FFFFFF", bold=True)
    metadata = [
        ("Gerado em (America/Sao_Paulo)", _generated_at(summary).strftime("%d/%m/%Y %H:%M:%S %Z (%z)")),
        ("Fuso horário", "America/Sao_Paulo"),
        ("Método", summary.get("method") or "Exportação do inventário informado; o gerador não consulta o banco, baixa imagens nem executa o pipeline."),
        ("Fonte", summary.get("source") or summary.get("source_path") or "Escopos e origem informados nas linhas do inventário."),
        ("Linhas exportadas", count),
        ("IDs de produtos distintos informados", len(product_ids)),
        ("IDs de imagens distintos informados", len(image_ids)),
        ("Linhas com erro informado", error_count),
        ("Estado anterior: Sim", before_counts["Sim"]),
        ("Estado anterior: Não", before_counts["Não"]),
        ("Estado anterior: Não verificado", before_counts[UNKNOWN]),
        ("Após execução: Sim", after_counts["Sim"]),
        ("Após execução: Não", after_counts["Não"]),
        ("Após execução: Não verificado", after_counts[UNKNOWN]),
        ("Como ler o estado do pipeline", SNAPSHOT_NOTE),
        ("URLs e recorte semântico", CROP_NOTE),
        ("Saídas locais", "Links de arquivos locais são relativos ao XLSX. Mantenha o relatório e a pasta de imagens juntos ao mover ou compartilhar o pacote."),
        ("Verificação visual", "O XLSX registra as evidências fornecidas. 'Sim' não substitui revisão visual da fidelidade, cor e ausência de fundo."),
    ]
    metadata.extend((f"Origem: {scope}", n) for scope, n in sorted(scope_counts.items()))
    metadata.extend((f"Processamento: {status}", n) for status, n in sorted(processing_counts.items()))
    metadata.extend((f"Dados adicionais: {key}", value) for key, value in summary.items()
                    if key not in {"generated_at", "method", "source", "source_path"})
    for row, (name, value) in enumerate(metadata, 2):
        _cell(overview, row, 1, name)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            overview.cell(row, 2, value)
        else:
            _cell(overview, row, 2, value)
        overview.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
        overview.row_dimensions[row].height = 70 if len(_text(value)) > 140 else 30

    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    return path
