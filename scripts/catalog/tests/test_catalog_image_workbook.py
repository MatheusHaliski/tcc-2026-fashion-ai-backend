"""Valida o artefato XLSX sem banco, rede ou download de imagens."""
import sys
from datetime import datetime, timezone
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from catalog_image_workbook import HEADERS, write_workbook  # noqa: E402


class CatalogImageWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root / "inventario.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def test_snapshot_unknown_status_and_hyperlinks_survive_reload(self):
        original = "https://cdn.example/shirt.jpg?view=front"
        processed = self.root / "imagens" / "shirt.png"
        processed.parent.mkdir()
        processed.write_bytes(b"local artifact")
        records = [{
            "product_id": "p1", "image_id": "i1", "product_name": "Camisa xadrez", "brand": "Marca",
            "source_url": original, "image_url": original, "standardized_before": None, "standardized_after": True,
            "processing_status": "APPROVED", "pipeline_version": "CATALOG_IMAGE_PIPELINE_V4",
            "crop_json": {"x": 0.2, "y": 0.1, "w": 0.5, "h": 0.4}, "output_path": processed,
            "status_note": "Recorte local concluído; snapshot anterior não comprova o banco.", "source_scope": "SNAPSHOT",
        }]
        result = write_workbook(iter(records), self.path, {"generated_at": datetime(2026, 10, 8, 18, 0, tzinfo=timezone.utc)})
        self.assertEqual(result, self.path)
        with ZipFile(self.path) as archive:
            self.assertIsNone(archive.testzip())
        book = load_workbook(self.path)
        try:
            sheet = book["Acervo"]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), HEADERS)
            self.assertEqual(sheet["D2"].value, "Não verificado")
            self.assertEqual(sheet["E2"].value, "Sim")
            self.assertEqual(sheet["C2"].hyperlink.target, original)
            self.assertEqual(sheet["F2"].hyperlink.target, original)
            self.assertEqual(sheet["G2"].hyperlink.target, "imagens/shirt.png")
            self.assertEqual(sheet.freeze_panes, "C2")
            self.assertEqual(sheet.auto_filter.ref, "A1:R2")
            self.assertEqual(sheet.tables["InventarioImagens"].ref, "A1:R2")
            overview = {row[0].value: row[1].value for row in book["Resumo"]}
            self.assertEqual(overview["Estado anterior: Não verificado"], 1)
            self.assertEqual(overview["Após execução: Sim"], 1)
            self.assertEqual(overview["Origem: SNAPSHOT"], 1)
            self.assertIn("15:00:00", overview["Gerado em (America/Sao_Paulo)"])
            self.assertIn("SEMANTIC_CROP", overview["URLs e recorte semântico"])
        finally:
            book.close()

    def test_untrusted_text_is_literal_and_credentialed_urls_are_omitted(self):
        rows = [{
            "product_name": "=HYPERLINK(\"https://evil.example\",\"clique\")", "brand": "+SUM(1,1)",
            "source_url": "https://user:private-password@cdn.example/a.jpg",
            "image_url": "https://cdn.example/a.jpg?access_token=private-token", "standardized_before": False,
            "standardized_after": False, "status_note": "@SUM(1,1)", "error": "URL https://user:private-password@cdn.example/a.jpg falhou.",
        }]
        write_workbook(rows, self.path, {"method": "=SUM(1,1)"})
        book = load_workbook(self.path, data_only=False)
        try:
            sheet = book["Acervo"]
            self.assertEqual(sheet["A2"].value, rows[0]["product_name"])
            self.assertEqual(sheet["A2"].data_type, "s")
            self.assertEqual(sheet["B2"].data_type, "s")
            self.assertIsNone(sheet["C2"].hyperlink)
            self.assertIsNone(sheet["F2"].hyperlink)
            for worksheet in book:
                self.assertFalse(any(cell.data_type == "f" for row in worksheet for cell in row))
        finally:
            book.close()
        with ZipFile(self.path) as archive:
            xml = b"".join(archive.read(name) for name in archive.namelist() if name.endswith(".xml") or name.endswith(".rels"))
            self.assertNotIn(b"private-password", xml)
            self.assertNotIn(b"private-token", xml)

    def test_summary_counts_do_not_coerce_missing_status_to_no(self):
        rows = [{"product_name": f"Peça {n}", "product_id": "p1", "image_id": str(n),
                 "standardized_before": before, "standardized_after": after,
                 "origin": "DATABASE", "processing_status": "APPROVED" if after else "PENDING"}
                for n, (before, after) in enumerate(((True, True), (False, True), (None, False)), 1)]
        write_workbook(rows, self.path)
        book = load_workbook(self.path)
        try:
            overview = {row[0].value: row[1].value for row in book["Resumo"]}
            self.assertEqual(overview["Linhas exportadas"], 3)
            self.assertEqual(overview["IDs de produtos distintos informados"], 1)
            self.assertEqual(overview["IDs de imagens distintos informados"], 3)
            self.assertEqual(overview["Estado anterior: Sim"], 1)
            self.assertEqual(overview["Estado anterior: Não"], 1)
            self.assertEqual(overview["Estado anterior: Não verificado"], 1)
            self.assertEqual(overview["Após execução: Sim"], 2)
            self.assertEqual(overview["Após execução: Não"], 1)
            self.assertEqual(overview["Origem: DATABASE"], 3)
        finally:
            book.close()

    def test_invalid_boolean_does_not_overwrite_previous_workbook(self):
        write_workbook([], self.path)
        before = self.path.read_bytes()
        with self.assertRaises(ValueError):
            write_workbook([{"standardized_before": "false"}], self.path)
        self.assertEqual(self.path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
