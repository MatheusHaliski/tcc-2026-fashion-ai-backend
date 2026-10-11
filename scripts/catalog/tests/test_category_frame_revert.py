"""Quadro antigo (V1/V2) desfeito quando a regra V3 recusa a foto (decisão REVERTED_TO_ORIGINAL).

A V3 recusa a foto (qualquer FRAME_UNAVAILABLE) e a imagem ainda mostra um JPEG do lote de uma versão anterior
(``pipeline_version`` V1/V2 e ``stored_url`` em /catalog/framed/): a referência ativa volta ao que estava antes do primeiro
quadro, pelo que a gravação antiga preservou (``previousStoredUrl``/``previousAssets``/``originalUrl``), pelo mesmo caminho
guardado e auditado do lote (guardas de concorrência, revisão humana, antes/depois, COMMIT incerto). Nada é apagado do S3.
Imagem que não mostra quadro antigo, revisão humana e execução sem --apply: nada muda.
"""
import io
import json
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stderr
from pathlib import Path
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import db
from catalog_image_updates import apply_product
from category_frame import LEGACY_VERSIONS, REVERTED_DECISION, VERSION, reverted_marker
from category_frame_storage import pre_frame_state, revert_plan
from test_catalog_image_updates import Database, Ranker, image
from test_category_frame_force import frame_of, response

from scripts.catalog.process_catalog_images import audit_record, main, missing_frame_environment

MEDIA = "https://media.example.com"
V1_URL = MEDIA + "/catalog/framed/1111/" + "a" * 64 + ".jpg"
V2_URL = MEDIA + "/catalog/framed/2222/" + "b" * 64 + ".jpg"
MASTER = "https://storage.example.com/masters/a.webp"
ORIGINAL = "https://example.com/original/a.jpg"
REFUSAL = "FRAME_UNAVAILABLE:COVER_FRAME_TOO_NARROW"
FRAME_ENV = {"MYSQL_HOST": "db.example", "MYSQL_DATABASE": "railway", "MYSQL_USER": "fai_app", "MYSQL_PASSWORD": "segredo-nunca-impresso",
             "S3_BUCKET": "bucket", "S3_ACCESS_KEY_ID": "chave", "S3_SECRET_ACCESS_KEY": "segredo-s3", "S3_SERVE_THROUGH_API": "false",
             "STORAGE_PUBLIC_BASE_URL": MEDIA}


def v1_assets(previous_url=MASTER):
    """Gravação V1 (#218/#220): sem previousAssets."""
    return {"card": V1_URL, "master": V1_URL, "framingVersion": "CATALOG_FRAME_34_50_V1", "sha256": "a" * 64,
            "previousStoredUrl": previous_url, "originalUrl": ORIGINAL}


def v2_assets(previous_url=MASTER, previous_assets=None):
    return {"card": V2_URL, "master": V2_URL, "framingVersion": "CATALOG_FRAME_34_FABRIC_V2", "sha256": "b" * 64,
            "width": 900, "height": 1200, "previousStoredUrl": previous_url,
            "previousAssets": previous_assets if previous_assets is not None else {"master": MASTER, "card": MASTER},
            "originalUrl": ORIGINAL, "editorFrame": {"version": "CATALOG_FRAME_34_FABRIC_V2", "policy": "FABRIC_ONLY_100"},
            "persistenceDecision": {"mode": "FORCED_CATEGORY_FRAME"}}


def legacy_image(image_id="a", assets=None, **changes):
    """Imagem mostrando o quadro V2 (close de tecido) gravado pelo lote antigo."""
    assets = assets if assets is not None else v2_assets()
    fields = dict(processing_status="APPROVED", review_status="NONE", pipeline_version="CATALOG_FRAME_34_FABRIC_V2",
                  is_canonical=True, view_role="CANONICAL", quality_score="0.8000", stored_url=assets["card"],
                  assets_json=json.dumps(assets),
                  crop_json=json.dumps({"aspect": "3:4", "crop": {"x": .4, "y": .3, "w": .1, "h": .1},
                                        "editorFrame": {"version": "CATALOG_FRAME_34_FABRIC_V2", "policy": "FABRIC_ONLY_100"}}),
                  metrics_json=json.dumps({"debug": {"editorFrame": {"version": "CATALOG_FRAME_34_FABRIC_V2"}}}))
    return image(image_id, **{**fields, **changes})


def refused_columns():
    """Colunas da análise Java atual da foto (nível A) que a V3 recusou."""
    return response(frame=frame_of(ok=False, reason="COVER_FRAME_TOO_NARROW"))["columns"]


def revert_analysis(record, reason=REFUSAL, columns="default"):
    plan = revert_plan(record, reason, product_frame_version="PRODUCT_RULE_FRAME_V2")
    return {"frame_revert": plan, "columns": refused_columns() if columns == "default" else columns}


class PreFrameStateTest(unittest.TestCase):
    def test_v2_volta_ao_asset_anterior_com_os_assets_preservados(self):
        state = pre_frame_state({"stored_url": V2_URL, "assets_json": v2_assets()})
        self.assertEqual((state["stored_url"], state["assets"]), (MASTER, {"master": MASTER, "card": MASTER}))
        self.assertEqual((state["restoredTo"], state["levels"], state["chainComplete"]), ("PREVIOUS_STORED_URL", 1, True))
        self.assertEqual(state["originalUrl"], ORIGINAL)

    def test_v2_sobre_v1_segue_a_cadeia_ate_antes_do_primeiro_quadro(self):
        # V2 gravada por cima da V1: o "anterior" da V2 é o JPEG da V1; o da V1 é o master de antes do lote
        stacked = v2_assets(previous_url=V1_URL, previous_assets=v1_assets())
        state = pre_frame_state({"stored_url": V2_URL, "assets_json": json.dumps(stacked)})
        self.assertEqual((state["stored_url"], state["levels"]), (MASTER, 2))
        self.assertIsNone(state["assets"])                       # a V1 não guardava previousAssets
        # V1 do #225 sobre snapshot: previousAssets em texto JSON
        text = {**v1_assets(), "previousAssets": json.dumps({"card": MASTER})}
        self.assertEqual(pre_frame_state({"stored_url": V1_URL, "assets_json": text})["assets"], {"card": MASTER})

    def test_sem_asset_anterior_ou_cadeia_incompleta_volta_a_foto_original_nunca_a_outro_quadro(self):
        never_stored = pre_frame_state({"stored_url": V2_URL, "assets_json": v2_assets(previous_url=None, previous_assets={})})
        self.assertEqual((never_stored["stored_url"], never_stored["assets"], never_stored["restoredTo"]), (None, None, "ORIGINAL_URL"))
        # quadro por cima de quadro sem o que estava antes do primeiro: foto original, sem assets
        broken = pre_frame_state({"stored_url": V2_URL, "assets_json": v2_assets(previous_url=V1_URL, previous_assets={})})
        self.assertEqual((broken["stored_url"], broken["assets"], broken["chainComplete"]), (None, None, False))
        # assets anteriores apontando para um quadro do lote não voltam
        odd = pre_frame_state({"stored_url": V2_URL, "assets_json": v2_assets(previous_assets={"card": V1_URL})})
        self.assertEqual((odd["stored_url"], odd["assets"]), (MASTER, None))

    def test_so_imagem_mostrando_quadro_antigo_tem_plano(self):
        self.assertIsNotNone(revert_plan({"pipeline_version": LEGACY_VERSIONS[0], "stored_url": V1_URL, "assets_json": v1_assets()}, REFUSAL))
        for record in ({"pipeline_version": "CATALOG_IMAGE_PIPELINE_V4", "stored_url": MASTER},
                       {"pipeline_version": VERSION, "stored_url": MEDIA + "/catalog/framed/3333/c.jpg"},
                       {"pipeline_version": "CATALOG_FRAME_34_FABRIC_V2", "stored_url": MASTER},
                       {"pipeline_version": "CATALOG_FRAME_34_FABRIC_V2", "stored_url": None}):
            with self.subTest(record=record):
                self.assertIsNone(revert_plan(record, REFUSAL))


class RevertWriteTest(unittest.TestCase):
    def database(self, *rows):
        conn = Database(list(rows))
        self.addCleanup(conn.sql.close)
        return conn

    def test_v2_recusada_pela_v3_volta_ao_estado_anterior_com_auditoria(self):
        conn = self.database(legacy_image())
        before = conn.rows()[0]
        records = conn.records()
        analysis = revert_analysis(records[0])
        result = apply_product(conn, records, {"a": analysis}, Ranker(conn), persistence_override="FORCED_CATEGORY_FRAME")
        row = conn.rows()[0]
        self.assertEqual(result["changed_ids"], ["a"])
        self.assertEqual(result["reverted_ids"], ["a"])
        # referência ativa: o master de antes do lote, com os assets preservados e o marcador da decisão V3
        self.assertEqual(row["stored_url"], MASTER)
        assets = json.loads(row["assets_json"])
        self.assertEqual((assets["master"], assets["card"]), (MASTER, MASTER))
        self.assertEqual(assets["editorFrame"]["decision"], REVERTED_DECISION)
        self.assertEqual(row["usage_status"], "PERSISTED")
        # o recorte 3:4 do close antigo não sobrevive: análise de nível A atual + marcador; versão da decisão V3
        crop = json.loads(row["crop_json"])
        self.assertEqual(crop["product"], {"x": .1, "y": .1, "w": .8, "h": .8})
        self.assertNotIn("aspect", crop)
        marker = crop["editorFrame"]
        self.assertEqual((marker["version"], marker["decision"], marker["reason"]), (VERSION, REVERTED_DECISION, REFUSAL))
        self.assertEqual((marker["revertedFrom"], marker["restoredTo"], marker["productFrameVersion"]),
                         ("CATALOG_FRAME_34_FABRIC_V2", "PREVIOUS_STORED_URL", "PRODUCT_RULE_FRAME_V2"))
        self.assertEqual(json.loads(row["metrics_json"])["debug"]["editorFrame"]["decision"], REVERTED_DECISION)
        self.assertEqual(row["pipeline_version"], VERSION)
        self.assertEqual(row["image_url"], before["image_url"])
        self.assertEqual(row["version"], before["version"] + 1)
        # auditoria: antes/depois da linha confirmada e o que foi desfeito
        change = result["changes"][0]
        self.assertEqual(change["before"]["stored_url"], V2_URL)
        self.assertEqual(change["after"]["stored_url"], MASTER)
        self.assertEqual(change["frame_revert"]["reason"], REFUSAL)
        self.assertEqual(change["frame_revert"]["framedUrl"], V2_URL)
        self.assertNotIn("persistence_decision_at_commit", change)    # restaurar não é persistir um asset novo
        self.assertFalse(any(sql.startswith(("INSERT", "DELETE")) for sql, _ in conn.statements))
        # a próxima execução reconhece a decisão e não refaz
        self.assertEqual(reverted_marker({**row, "crop_json": row["crop_json"]})["decision"], REVERTED_DECISION)

    def test_sem_asset_anterior_volta_a_foto_original_e_sem_analise_java_limpa_o_recorte_antigo(self):
        assets = v2_assets(previous_url=None, previous_assets={})
        conn = self.database(legacy_image(assets=assets))
        records = conn.records()
        analysis = revert_analysis(records[0], reason="FRAME_UNAVAILABLE:JAVA_IMAGE_TOO_SMALL", columns=None)
        apply_product(conn, records, {"a": analysis}, Ranker(conn), persistence_override="FORCED_CATEGORY_FRAME")
        row = conn.rows()[0]
        self.assertIsNone(row["stored_url"])
        self.assertIsNone(row["assets_json"])                  # o card volta à foto original (image_url)
        self.assertEqual(row["usage_status"], "REFERENCE_ONLY")
        self.assertIsNone(row["crop_json"])                    # nem o recorte do close antigo aplicado à original
        self.assertEqual(json.loads(row["metrics_json"])["debug"]["editorFrame"]["restoredTo"], "ORIGINAL_URL")
        self.assertEqual(row["pipeline_version"], VERSION)

    def test_imagem_que_nao_mostra_quadro_antigo_fica_intocada(self):
        conn = self.database(image("a", processing_status="APPROVED", pipeline_version="CATALOG_IMAGE_PIPELINE_V4"))
        before = conn.rows()
        records = conn.records()
        # plano feito sobre um estado antigo: no COMMIT a imagem não é mais V1/V2 enquadrada
        stale = revert_analysis({**records[0], "pipeline_version": "CATALOG_FRAME_34_FABRIC_V2", "stored_url": V2_URL,
                                 "assets_json": v2_assets()})
        result = apply_product(conn, records, {"a": stale}, Ranker(conn))
        self.assertEqual(conn.rows(), before)
        self.assertEqual(result["skip_reasons"], {"a": "NOT_LEGACY_FRAME"})
        self.assertFalse(conn.statements)

    def test_revisao_humana_fica_intocada(self):
        for protected in ({"review_status": "APPROVED"}, {"review_status": "REJECTED"}, {"usage_status": "REJECTED"},
                          {"processing_status": "DOWNLOADING"}):
            with self.subTest(protected=protected):
                conn = self.database(legacy_image(**protected))
                before = conn.rows()
                records = conn.records()
                result = apply_product(conn, records, {"a": revert_analysis(records[0])}, Ranker(conn))
                self.assertEqual(conn.rows(), before)
                self.assertFalse(result["changed_ids"])
                self.assertFalse(conn.statements)

    def test_referencia_ativa_mudada_depois_do_inventario_pula_o_produto(self):
        conn = self.database(legacy_image(), image("b"))
        records = conn.records()
        # outra gravação trocou o asset ativo sem passar pela versão: a restauração não pode desfazer o que não leu
        conn.sql.execute("UPDATE catalog_images SET stored_url=? WHERE id='a'", (MEDIA + "/catalog/framed/9999/x.jpg",))
        conn.sql.commit()
        before = conn.rows()
        result = apply_product(conn, records, {"a": revert_analysis(records[0])}, Ranker(conn))
        self.assertEqual(conn.rows(), before)
        self.assertEqual(result["skip_reasons"], {"a": "PRODUCT_CHANGED"})

    def test_plano_que_restauraria_outro_quadro_e_recusado_antes_do_sql(self):
        conn = self.database(legacy_image())
        records = conn.records()
        analysis = revert_analysis(records[0])
        analysis["frame_revert"]["stored_url"] = V1_URL
        with self.assertRaisesRegex(ValueError, "another framed asset"):
            apply_product(conn, records, {"a": analysis}, Ranker(conn))
        self.assertFalse(conn.statements)

    @unittest.skipIf(db.pymysql is None, "PyMySQL required to simulate transport errors")
    def test_commit_incerto_nao_duplica_versao_nem_auditoria(self):
        for fault in ("before", "after"):
            with self.subTest(fault=fault):
                conn = self.database(legacy_image())
                conn.drop_commit = fault
                records = conn.records()
                with patch("db.time.sleep"):
                    result = apply_product(conn, records, {"a": revert_analysis(records[0])}, Ranker(conn))
                self.assertEqual(len(result["changes"]), 1)
                self.assertEqual(conn.rows()[0]["version"], 5)
                self.assertEqual(conn.rows()[0]["stored_url"], MASTER)
                self.assertEqual(result.get("recovered_commit", False), fault == "after")


class RevertAuditRecordTest(unittest.TestCase):
    def analyzer(self, frame=None):
        analyzer = Mock(spec=["analyze", "rank", "ready", "category_frame", "force_frame", "frame_storage"])
        analyzer.ready = {"pipelineVersion": "CATALOG_IMAGE_PIPELINE_V4", "productFrameVersion": "PRODUCT_RULE_FRAME_V2"}
        analyzer.category_frame = True
        analyzer.force_frame = True
        analyzer.frame_storage = Mock()
        analyzer.analyze.return_value = response(frame=frame or frame_of(ok=False, reason="COVER_FRAME_TOO_NARROW"))
        return analyzer

    def run_record(self, record, analyzer):
        downloader, checkpoint = Mock(), Mock()
        checkpoint.get.return_value = None
        downloader.get.return_value = Path("/tmp/a.img")
        downloader.final_urls = {}
        return audit_record(record, downloader, analyzer, checkpoint, apply=True), downloader

    def record(self, **changes):
        row = {**legacy_image(), "image_id": "a", "category": "upper_piece", "source_url": ORIGINAL,
               "allows_image_persistence": False, "persistence_decision": {"state": "UNKNOWN", "reason": "SOURCE_MISSING", "domains": []}}
        row.update(changes)
        return row

    def test_recusa_da_v3_em_imagem_v2_planeja_a_restauracao_com_a_analise_atual(self):
        analyzer = self.analyzer()
        result, _ = self.run_record(self.record(), analyzer)
        self.assertEqual(result["error"], REFUSAL)
        self.assertEqual(result["frame_revert"]["stored_url"], MASTER)
        self.assertEqual(result["analysis"]["frame_revert"], result["frame_revert"])
        self.assertEqual(result["analysis"]["columns"]["pipeline_version"], "CATALOG_IMAGE_PIPELINE_V4")
        self.assertIn("volta ao estado anterior", result["status_note"])
        analyzer.frame_storage.store.assert_not_called()          # nada é enviado ao S3

    def test_recusa_em_imagem_sem_quadro_antigo_nao_planeja_nada(self):
        for changes in ({"pipeline_version": "CATALOG_IMAGE_PIPELINE_V4", "stored_url": None, "assets_json": None},
                        {"pipeline_version": "CATALOG_FRAME_34_FABRIC_V2", "stored_url": MASTER}):
            with self.subTest(changes=changes):
                result, _ = self.run_record(self.record(**changes), self.analyzer())
                self.assertEqual(result["error"], REFUSAL)
                self.assertNotIn("frame_revert", result)
                self.assertNotIn("analysis", result)

    def test_revisao_humana_nao_e_analisada_nem_restaurada(self):
        result, downloader = self.run_record(self.record(review_status="APPROVED"), self.analyzer())
        self.assertIn("Preservado", result["status_note"])
        self.assertNotIn("frame_revert", result)
        downloader.get.assert_not_called()

    def test_quadro_aceito_pela_v3_nao_e_restauracao(self):
        storage = Mock()
        storage.store.return_value = {"stored_url": MEDIA + "/catalog/framed/4444/d.jpg", "assets_json": "{}", "owned_key": "k", "render_info": {}}
        analyzer = self.analyzer(frame=frame_of())
        analyzer.frame_storage = storage
        with patch("scripts.catalog.process_catalog_images.public_image_url", side_effect=lambda url: url):
            result, _ = self.run_record(self.record(), analyzer)
        self.assertNotIn("error", result)
        self.assertNotIn("frame_revert", result)
        self.assertIn("framed_assets", result["analysis"])

    def test_imagem_ja_restaurada_pela_mesma_regra_nao_e_baixada_de_novo(self):
        marker = {"version": VERSION, "decision": REVERTED_DECISION, "reason": REFUSAL, "productFrameVersion": "PRODUCT_RULE_FRAME_V2"}
        done = self.record(pipeline_version=VERSION, stored_url=MASTER, assets_json=json.dumps({"card": MASTER, "editorFrame": marker}),
                           crop_json=json.dumps({"aspect": "4:5", "crop": {"x": .1, "y": .1, "w": .8, "h": .8}, "editorFrame": marker}))
        result, downloader = self.run_record(done, self.analyzer())
        self.assertIn("já foi desfeito", result["status_note"])
        downloader.get.assert_not_called()
        # outra versão da regra: analisa de novo (pode ganhar quadro agora)
        marker["productFrameVersion"] = "PRODUCT_RULE_FRAME_V1"
        again = self.record(pipeline_version=VERSION, stored_url=MASTER, crop_json=json.dumps({"editorFrame": marker}))
        _, downloader = self.run_record(again, self.analyzer())
        downloader.get.assert_called()


class RevertCliTest(unittest.TestCase):
    """main ponta a ponta com o MySQL simulado (SQLite) e o Java/S3 substituídos."""

    def run_main(self, conn, argv, analyzer_frame):
        records = conn.records()
        for record in records:
            record.update(allows_image_persistence=False, persistence_decision={"state": "UNKNOWN", "reason": "SOURCE_MISSING", "domains": []})
        ranker = Ranker(conn)

        class Analyzer:
            def __init__(self, *args, **kwargs):
                self.ready = {"pipelineVersion": "CATALOG_IMAGE_PIPELINE_V4", "productFrameVersion": "PRODUCT_RULE_FRAME_V2"}

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def analyze(self, *args, **kwargs):
                return response(frame=analyzer_frame)

            def rank(self, category, candidates):
                return ranker.rank(category, candidates)

        downloader = Mock()
        downloader.get.return_value = Path("/tmp/a.img")
        downloader.final_urls = {}
        downloader.proxy_blocked = False
        storage = Mock()
        storage.pending_assets.return_value = []
        connect = Mock(return_value=conn)
        apply_spy = Mock(wraps=apply_product)
        with patch.dict(os.environ, FRAME_ENV), \
                patch("catalog_image_inventory.load_database", return_value=records), \
                patch("catalog_image_analyzer.CatalogImageAnalyzer", Analyzer), \
                patch("catalog_image_analyzer.ensure_classpath", return_value="cp"), \
                patch("category_frame_storage.FrameStorage", return_value=storage), \
                patch("scripts.catalog.process_catalog_images.ImageDownloader", return_value=downloader), \
                patch("catalog_image_updates.apply_product", apply_spy), \
                patch("db.connect", connect), redirect_stderr(io.StringIO()) as err, patch("sys.stdout", new=io.StringIO()):
            code = main(argv)
        return code, err.getvalue(), connect, apply_spy, storage

    def test_lote_com_apply_restaura_conta_no_resumo_na_planilha_e_audita(self):
        from openpyxl import load_workbook
        conn = Database([legacy_image()])
        self.addCleanup(conn.sql.close)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "lote.xlsx"
            code, err, _, _, storage = self.run_main(conn, ["--database", "--apply", "--category-frame", "--force-category-frame",
                                                            "--output", str(output)], frame_of(ok=False, reason="COVER_FRAME_TOO_NARROW"))
            self.assertEqual(code, 2, err)                       # a foto continua sem quadro V3: falha individual registrada
            row = conn.rows()[0]
            self.assertEqual((row["stored_url"], row["pipeline_version"]), (MASTER, VERSION))
            summary = json.loads(output.with_suffix(".summary.json").read_text())
            self.assertEqual(summary["frame_reverted"], 1)
            self.assertEqual(summary["frame_unavailable"], {"COVER_FRAME_TOO_NARROW": 1})
            self.assertEqual(summary["frame_observations"], {"REVERTED_TO_ORIGINAL": 1})
            journal = [json.loads(line) for line in output.with_suffix(".changes.audit.jsonl").read_text().splitlines()]
            self.assertEqual(len(journal), 1)
            self.assertEqual(journal[0]["frame_revert"]["reason"], REFUSAL)
            self.assertIn("url_sha256", json.dumps(journal[0]["before"]["stored_url"]))   # URLs só como hash na auditoria
            self.assertIn("url_sha256", json.dumps(journal[0]["after"]["stored_url"]))
            book = load_workbook(output)
            overview = {r[0].value: r[1].value for r in book["Resumo"].iter_rows(min_row=2)}
            self.assertEqual(overview["Quadros antigos (V1/V2) desfeitos: V3 recusou a foto"], 1)
            notes = [r[8].value for r in book["Acervo"].iter_rows(min_row=2)]
            self.assertIn("voltou ao estado anterior ao quadro", notes[0])
            storage.store.assert_not_called()
            storage.discard_unreferenced.assert_not_called()    # nenhum objeto do S3 é apagado

    def test_sem_apply_nada_e_gravado(self):
        conn = Database([legacy_image()])
        self.addCleanup(conn.sql.close)
        before = conn.rows()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "auditoria.xlsx"
            code, err, connect, apply_spy, storage = self.run_main(conn, ["--database", "--output", str(output)],
                                                                   frame_of(ok=False, reason="COVER_FRAME_TOO_NARROW"))
            self.assertEqual(code, 0, err)
            summary = json.loads(output.with_suffix(".summary.json").read_text())
            self.assertEqual(summary["frame_reverted"], 0)
            self.assertFalse(output.with_suffix(".changes.audit.jsonl").exists())
        self.assertEqual(conn.rows(), before)
        self.assertFalse(conn.statements)
        connect.assert_not_called()
        apply_spy.assert_not_called()


class FrameEnvironmentTest(unittest.TestCase):
    def test_falta_de_variaveis_para_antes_de_ler_o_banco_ou_iniciar_o_java_sem_exibir_valores(self):
        env = {k: v for k, v in FRAME_ENV.items() if k not in ("S3_BUCKET", "S3_SECRET_ACCESS_KEY", "MYSQL_HOST")}
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, env, clear=True), \
                patch("catalog_image_inventory.load_database") as load, \
                patch("catalog_image_analyzer.ensure_classpath") as classpath, redirect_stderr(io.StringIO()) as err:
            code = main(["--database", "--apply", "--category-frame", "--force-category-frame", "--output", str(Path(directory) / "o.xlsx")])
        self.assertEqual(code, 2)
        message = err.getvalue()
        for name in ("S3_BUCKET", "S3_SECRET_ACCESS_KEY", "MYSQL_HOST"):
            self.assertIn(name, message)
        for value in FRAME_ENV.values():
            if len(value) > 5:
                self.assertNotIn(value, message)
        self.assertNotIn("Traceback", message)
        load.assert_not_called()
        classpath.assert_not_called()

    def test_url_publica_do_bucket_sem_https_para_antes_do_java_sem_traceback(self):
        env = {**FRAME_ENV, "STORAGE_PUBLIC_BASE_URL": "http://media.example.com"}
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, env, clear=True), \
                patch("catalog_image_inventory.load_database", return_value=[]), \
                patch("catalog_image_analyzer.ensure_classpath") as classpath, redirect_stderr(io.StringIO()) as err:
            code = main(["--database", "--apply", "--category-frame", "--output", str(Path(directory) / "o.xlsx")])
        self.assertEqual(code, 2)
        self.assertIn("STORAGE_PUBLIC_BASE_URL", err.getvalue())
        self.assertNotIn("Traceback", err.getvalue())
        self.assertNotIn("segredo", err.getvalue())
        classpath.assert_not_called()

    def test_nomes_exigidos_e_url_publica_do_bucket(self):
        self.assertEqual(missing_frame_environment(FRAME_ENV), [])
        self.assertEqual(missing_frame_environment({**FRAME_ENV, "S3_BUCKET": " "}), ["S3_BUCKET"])
        through_api = {**FRAME_ENV, "S3_SERVE_THROUGH_API": "true", "STORAGE_PUBLIC_BASE_URL": ""}
        self.assertEqual(missing_frame_environment(through_api), ["APP_BASE_URL"])
        self.assertEqual(missing_frame_environment({**through_api, "APP_BASE_URL": "https://api.example.com"}), [])
        self.assertIn("STORAGE_PUBLIC_BASE_URL (ou APP_BASE_URL com S3_SERVE_THROUGH_API=true)",
                      missing_frame_environment({**FRAME_ENV, "STORAGE_PUBLIC_BASE_URL": ""}))
        self.assertEqual(missing_frame_environment({}), ["MYSQL_HOST", "MYSQL_DATABASE", "MYSQL_USER", "MYSQL_PASSWORD", "S3_BUCKET",
                                                     "S3_ACCESS_KEY_ID", "S3_SECRET_ACCESS_KEY",
                                                     "STORAGE_PUBLIC_BASE_URL (ou APP_BASE_URL com S3_SERVE_THROUGH_API=true)"])


if __name__ == "__main__":
    unittest.main()
