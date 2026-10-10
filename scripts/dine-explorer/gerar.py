# -*- coding: utf-8 -*-
"""Gera os diagramas UML (atividades, classes, máquina de estados, sequência e
componentes), a planilha de entidades × bancos e o JSON do Trello do Dine
Explorer a partir de dados_rf.py.

    python3 scripts/dine-explorer/gerar.py                 # .puml + .xlsx + .json + índice
    python3 scripts/dine-explorer/gerar.py --render JAR    # idem + PNG (java -jar plantuml.jar)
"""
import json
import os
import re
import subprocess
import sys

sys.path.insert(0, os.path.dirname(__file__))
from dados_rf import RFS, N, SPRING  # noqa: E402

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS = os.path.join(RAIZ, "docs", "dine-explorer")
DIAG = os.path.join(DOCS, "diagramas")

DB_ALIAS = {"MySQL": "DB_MySQL", "Cassandra": "DB_Cassandra", "Redis": "DB_Redis", "OpenSearch": "DB_OpenSearch", "S3": "DB_S3"}
DB_LABEL = {"MySQL": "MySQL 8.4\\n(Spring Data JPA + Flyway)", "Cassandra": "Cassandra 5\\n(Spring Data Cassandra)",
            "Redis": "Redis 7\\n(Spring Data Redis)", "OpenSearch": "OpenSearch 2\\n(opensearch-java)", "S3": "Amazon S3\\n(AWS SDK v2)"}

BASE = """skinparam backgroundColor #FFFFFF
skinparam shadowing false
skinparam defaultFontName "DejaVu Sans"
skinparam defaultFontSize 12
skinparam titleFontSize 15
skinparam roundcorner 14
skinparam ArrowColor #52606D
skinparam note {
  BackgroundColor #FFFDEB
  BorderColor #D8CE78
  FontColor #554F24
}"""


def esc(s):
    return s.replace('"', "'")


def titulo(r, tipo):
    t = r["titulo"]
    if len(t) > 110:
        t = t[:107] + "…"
    return f"title RF{r['n']} — {t}\\nDiagrama de {tipo} — Dine Explorer · {SPRING}"


def wrap(s, n=38):
    out, linha = [], ""
    for p in s.split():
        if len(linha) + len(p) + 1 > n and linha:
            out.append(linha)
            linha = p
        else:
            linha = (linha + " " + p).strip()
    if linha:
        out.append(linha)
    return "\\n".join(out)


def tabela(e):
    return e[1].split(" ")[0]


def dbs(r):
    seen = []
    for e in r["entidades"]:
        if e[2] not in seen:
            seen.append(e[2])
    return seen


# ------------------------------------------------------------------ atividades
def atividades(r):
    L = [f"@startuml RF{r['n']}_Atividades", BASE, """skinparam activity {
  BackgroundColor #F4F2EF
  BorderColor #4A4852
  DiamondBackgroundColor #FFF4E5
  DiamondBorderColor #C9822B
  StartColor #20A77A
  EndColor #4B5563
}
skinparam swimlane {
  BorderColor #CBC6BE
  TitleFontColor #2D2438
}""", titulo(r, "atividades"), ""]
    L += ["|Pessoa usuária|", "start", f":Abre {r['tela']};", "|App (Next.js 15 / React Native)|"]
    gets = [x for x in r["rotas"] if x.startswith("GET")]
    if gets:
        L.append(":Carrega os dados da tela\\n" + "\\n".join(gets[:3]) + ";")
    L.append("|Pessoa usuária|")
    for i, t in enumerate(r["telas"]):
        if i == 0:
            L.append(f":{wrap(t)};")
    L.append("|App (Next.js 15 / React Native)|")
    acao = next((x for x in r["rotas"] if not x.startswith("GET")), r["rotas"][0])
    L.append(f":Valida no cliente (Zod)\\n{acao};")
    L.append(f"|API ({SPRING.split(' ·')[0]})|")
    L.append(f":{r['controller']}\\n→ {r['service']}.{r['metodos'][0]}\\n(Bean Validation + Spring Security 7);")
    for erro in r["erros"][:4]:
        L.append(f"if ({erro.lower().replace('_', ' ')}?) then (sim)")
        L.append(f"  :4xx {erro}\\n(ProblemDetail RFC 9457)\\n→ a tela mostra o aviso;")
        L.append("  stop")
        L.append("endif")
    if len(r["erros"]) > 4:
        L.append("note right\n  e mais: " + ", ".join(r["erros"][4:]) + "\nend note")
    L.append("|Dados e serviços|")
    for db in dbs(r):
        ents = [tabela(e) for e in r["entidades"] if e[2] == db]
        L.append(f":{db}: " + wrap(", ".join(ents), 40) + ";")
    for p in r["portas"][:3]:
        L.append(f":{p.split(' «')[0]};")
    L.append(f":Publica evento de domínio\\n(Spring Modulith → Kafka)\\n→ métricas RF24 / auditoria;")
    L.append("|App (Next.js 15 / React Native)|")
    L.append(":Atualiza a tela (TanStack Query)\\ne mostra feedback;")
    L.append("|Pessoa usuária|")
    if len(r["telas"]) > 1:
        L.append("if (Faz outra ação?) then (sim)")
        L.append("  switch (Qual?)")
        for t, rota in zip(r["telas"][1:], (r["rotas"][1:] + r["rotas"])[: len(r["telas"]) - 1]):
            L.append(f"  case ({esc(t[:34])}{'…' if len(t) > 34 else ''})")
            L.append(f"    :{wrap(t, 40)}\\n→ {rota};")
        L.append("  endswitch")
        L.append("endif")
    L.append("stop")
    L.append("@enduml")
    return "\n".join(L)


# ------------------------------------------------------------------ sequência
def sequencia(r):
    L = [f"@startuml RF{r['n']}_Sequencia", BASE, """skinparam sequence {
  ArrowColor #4A4852
  LifeLineBorderColor #CBC6BE
  ParticipantBackgroundColor #F4F2EF
  ParticipantBorderColor #4A4852
  ActorBorderColor #4A4852
  GroupBorderColor #CBC6BE
  GroupBackgroundColor #FAF8F5
}
autonumber "<b>0."
""", titulo(r, "sequência"), ""]
    L += ['actor "Pessoa usuária" as U', 'boundary "App (Next.js / RN)" as P', 'participant "API Gateway\\n(Spring Cloud Gateway)" as GW',
          f'control "{r["controller"]}" as C', f'participant "{r["service"]}" as S']
    portas = r["portas"][:4]
    for i, p in enumerate(portas):
        nome, _, impl = p.partition(" «")
        L.append(f'participant "{nome}\\n({impl.rstrip("»")})" as P{i}' if impl else f'participant "{nome}" as P{i}')
    for db in dbs(r):
        L.append(f'database "{DB_LABEL[db]}" as {DB_ALIAS[db]}')
    L.append('queue "Kafka" as K')
    L.append(f"== {esc(r['telas'][0][:80])} ==")
    acao = next((x for x in r["rotas"] if not x.startswith("GET")), r["rotas"][0])
    L += ["U -> P : interage na tela", f"P -> GW : {acao}\\nAuthorization: Bearer JWT",
          "GW -> GW : valida JWT (ES256) + rate limit (Redis)", f"GW -> C : {acao}", f"C -> S : {r['metodos'][0]}", "activate S"]
    if "Redis" in dbs(r):
        L.append("S -> DB_Redis : GET cache / lock")
    if "MySQL" in dbs(r):
        ms = [tabela(e) for e in r["entidades"] if e[2] == "MySQL"][:3]
        L.append(f"S -> DB_MySQL : SELECT {', '.join(ms)}")
    for i, e in enumerate(r["erros"][:2]):
        L += [f"alt {e.lower().replace('_', ' ')}", f"S --> C : {e}", f"C --> P : 4xx ProblemDetail {{code: {e}}}", "P --> U : mostra aviso", "end"]
    for i, _ in enumerate(portas):
        L.append(f"S -> P{i} : {r['metodos'][min(i, len(r['metodos']) - 1)]}")
        L.append(f"P{i} --> S : ok")
    if "MySQL" in dbs(r):
        L.append("S -> DB_MySQL : INSERT/UPDATE (@Transactional)")
    for db in ("Cassandra", "OpenSearch", "S3"):
        if db in dbs(r):
            ents = [tabela(e) for e in r["entidades"] if e[2] == db][:2]
            verbo = {"Cassandra": "INSERT", "OpenSearch": "index/search", "S3": "PUT/GET presigned"}[db]
            L.append(f"S -> {DB_ALIAS[db]} : {verbo} {', '.join(ents)}")
    if "Redis" in dbs(r):
        L.append("S -> DB_Redis : SET/ZADD (invalida cache)")
    L += ["S ->> K : evento de domínio (outbox)", "deactivate S", "S --> C : DTO (record)", "C --> P : 200/201 JSON", "P --> U : atualiza a tela"]
    for t, rota in zip(r["telas"][1:3], r["rotas"][1:3]):
        L += [f"== {esc(t[:80])} ==", f"U -> P : {esc(t[:50])}", f"P -> C : {rota}", f"C -> S : {r['metodos'][min(1, len(r['metodos']) - 1)]}",
              "S --> C : resultado", "C --> P : 200 JSON"]
    L.append("@enduml")
    return "\n".join(L)


# ------------------------------------------------------------------ classes
def classes(r):
    L = [f"@startuml RF{r['n']}_Classes", "!pragma layout smetana", BASE, """skinparam class {
  BackgroundColor #F4F2EF
  BorderColor #4A4852
  ArrowColor #4A4852
  StereotypeFontColor #6B6470
}
hide empty members
left to right direction""", titulo(r, "classes"), ""]
    porbanco = {}
    for e in r["entidades"]:
        porbanco.setdefault(e[2], []).append(e)
    nomes = {e[0] for e in r["entidades"]}
    for db, ents in porbanco.items():
        L.append(f'package "{db}" {{')
        for e in ents:
            L.append(f'  class {e[0]} <<{tabela(e)}>> {{')
            for a in e[3]:
                L.append(f"    +{a}")
            L.append("  }")
        L.append("}")
    L.append('package "dominio (enums)" {')
    for en, vals in r["enums"].items():
        L.append(f"  enum {en} {{")
        for v in vals:
            L.append(f"    {v}")
        L.append("  }")
    L.append("}")
    L.append('package "aplicacao (casos de uso)" {')
    L.append(f"  class {r['service']} <<@Service>> {{")
    for m in r["metodos"]:
        L.append(f"    +{m}")
    L.append("  }")
    L.append(f"  class {r['controller']} <<@RestController>> {{")
    for rota in r["rotas"][:6]:
        L.append(f"    +{rota.split(' (')[0]}")
    L.append("  }")
    for p in r["portas"]:
        nome = re.sub(r"\W", "", p.split(" «")[0].split(" (")[0])
        if nome and nome not in nomes:
            L.append(f"  interface {nome} <<port>>")
            nomes.add(nome)
            L.append(f"  {r['service']} ..> {nome}")
    L.append("}")
    L.append(f"{r['controller']} --> {r['service']}")
    for e in r["entidades"]:
        L.append(f"{r['service']} ..> {e[0]} : usa")
    # relações por atributo tipado com outra entidade
    ent_names = {e[0] for e in r["entidades"]}
    for e in r["entidades"]:
        for a in e[3]:
            m = re.match(r"(\w+): (List<|Set<)?(\w+)", a)
            if m and m.group(3) in ent_names and m.group(3) != e[0]:
                mult = '"*"' if m.group(2) else '"1"'
                L.append(f'{e[0]} "*" --> {mult} {m.group(3)} : {m.group(1)}')
        for en in r["enums"]:
            if any(en in a for a in e[3]):
                L.append(f"{e[0]} --> {en}")
    L.append("@enduml")
    return "\n".join(L)


# ------------------------------------------------------------------ estados
def estados(r):
    s = r["estado"]
    L = [f"@startuml RF{r['n']}_MaquinaDeEstados_{s['nome']}", "!pragma layout smetana", titulo(r, f"máquina de estados — {s['nome']}"),
         """skinparam backgroundColor #FFFFFF
skinparam state {
  BackgroundColor #F8F4FF
  BorderColor #B695F5
  FontColor #2D2438
  ArrowColor #52606D
}
skinparam note {
  BackgroundColor #FFFDEB
  BorderColor #D8CE78
  FontColor #554F24
}"""]
    for a, rot in s["estados"]:
        L.append(f'state "{rot.replace(chr(10), chr(92) + "n")}" as {a}')
    L.append("")
    for de, para, ev in s["transicoes"]:
        ev = ev.replace("\n", "\\n")
        L.append(f"{de} --> {para}" + (f" : {ev}" if ev else ""))
    if s.get("nota"):
        L += [f"note right of {s['estados'][0][0]}", "  " + s["nota"].replace("\n", "\n  "), "end note"]
    L.append("@enduml")
    return "\n".join(L)


# ------------------------------------------------------------------ componentes
def componentes(r):
    L = [f"@startuml RF{r['n']}_Componentes", "!pragma layout smetana", BASE, """skinparam component {
  BackgroundColor #F4F2EF
  BorderColor #4A4852
}
skinparam package {
  BorderColor #CBC6BE
  FontColor #4A4852
}
skinparam database {
  BackgroundColor #EEF6FF
  BorderColor #4A4852
}
skinparam cloud {
  BackgroundColor #F0FBF6
  BorderColor #4A4852
}
left to right direction""", titulo(r, "componentes"), ""]
    L.append('package "Frontend (Next.js 15 · React 19 · TanStack Query)" {')
    L.append(f'  component "{esc(r["tela"])}" as FE')
    if r["n"] in (27, 28):
        L.append('  component "GlobeView\\n(React Three Fiber + three-globe\\n+ deck.gl H3HexagonLayer + Lottie/Rive)" as FE3D')
    L.append('  component "lib/api (OpenAPI client gerado)" as APIC')
    L.append("}")
    L.append('package "Borda" {\n  component "CloudFront + WAF" as CDN\n  component "Spring Cloud Gateway\\n(JWT, rate limit, CORS)" as GW\n}')
    L.append(f'package "dine-web (REST · {SPRING.split(" ·")[0]})" {{')
    L.append(f'  component "{r["controller"]}\\n' + "\\n".join(x.split(" (")[0] for x in r["rotas"][:5]) + (f"\\n(+{len(r['rotas']) - 5} rotas)" if len(r["rotas"]) > 5 else "") + '" as CT')
    L.append("}")
    L.append('package "dine-application (Spring Modulith · casos de uso)" {')
    L.append(f'  component "{r["service"]}" as SV')
    for i, p in enumerate(r["portas"]):
        nome, _, impl = p.partition(" «")
        L.append(f'  component "{nome}' + (f'\\n«{impl.rstrip("»")}»' if impl else "") + f'" as PT{i} <<port>>')
    L.append("}")
    L.append('package "dine-domain (entidades + repositórios)" {')
    for db in dbs(r):
        ents = [e[0] + "Repository" for e in r["entidades"] if e[2] == db and db != "S3"]
        if ents:
            L.append(f'  component "' + "\\n".join(ents) + f'" as REPO_{db}')
    L.append("}")
    for db in dbs(r):
        tabs = [tabela(e) for e in r["entidades"] if e[2] == db]
        L.append(f'database "{DB_LABEL[db]}\\n• ' + "\\n• ".join(tabs) + f'" as {DB_ALIAS[db]}')
    for i, x in enumerate(r["externos"]):
        L.append(f'cloud "{x}" as EX{i}')
    L.append('component "Observabilidade\\n(Micrometer + OpenTelemetry\\n→ Grafana/Tempo/Loki)" as OBS')
    L.append("FE --> APIC")
    if r["n"] in (27, 28):
        L.append("FE --> FE3D")
        L.append("FE3D --> CDN : tiles H3 / Lottie")
    L += ["APIC --> CDN : HTTPS/JSON", "CDN --> GW", "GW --> CT", "CT --> SV", "SV ..> OBS"]
    for i, _ in enumerate(r["portas"]):
        L.append(f"SV --> PT{i}")
    for db in dbs(r):
        if db == "S3":
            continue
        L.append(f"SV --> REPO_{db}")
        L.append(f"REPO_{db} --> {DB_ALIAS[db]}")
    if "S3" in dbs(r):
        s3p = next((i for i, p in enumerate(r["portas"]) if "S3" in p or "Storage" in p), None)
        L.append(f"{'PT' + str(s3p) if s3p is not None else 'SV'} --> DB_S3")
    for i, x in enumerate(r["externos"]):
        alvo = next((j for j, p in enumerate(r["portas"]) if any(w.lower() in p.lower() for w in x.split()[:2] if len(w) > 2)), None)
        L.append(f"{'PT' + str(alvo) if alvo is not None else 'SV'} --> EX{i}")
    L.append("@enduml")
    return "\n".join(L)


TIPOS = [("atividades", atividades, "Atividades"), ("sequencia", sequencia, "Sequencia"), ("componentes", componentes, "Componentes"),
         ("maquinadeestados", estados, "MaquinaDeEstados"), ("classes", classes, "Classes")]


def gerar_diagramas():
    arquivos = []
    for r in RFS:
        d = os.path.join(DIAG, f"RF{r['n']}")
        os.makedirs(d, exist_ok=True)
        for slug, fn, _ in TIPOS:
            nome = f"RF{r['n']}-{slug}" + (f"-{r['estado']['nome'].lower()}" if slug == "maquinadeestados" else "") + ".puml"
            p = os.path.join(d, nome)
            with open(p, "w", encoding="utf-8") as f:
                f.write(fn(r) + "\n")
            arquivos.append(p)
    return arquivos


def render(jar, arquivos):
    for i in range(0, len(arquivos), 25):
        subprocess.run(["java", "-Djava.awt.headless=true", "-DPLANTUML_LIMIT_SIZE=16384", "-jar", jar, "-tpng", "-charset", "UTF-8", *arquivos[i:i + 25]], check=False)


def indice():
    L = ["# Dine Explorer — índice dos diagramas UML", "",
         f"Gerado por `scripts/dine-explorer/gerar.py` a partir de `scripts/dine-explorer/dados_rf.py` (RF1–RF{N}).",
         f"Arquitetura: **{SPRING}**, Spring Modulith, MySQL · Cassandra · Redis · OpenSearch · S3.", "",
         "```bash", "python3 scripts/dine-explorer/gerar.py --render /caminho/plantuml.jar", "```", "",
         "| RF | Funcionalidade | Atividades | Sequência | Componentes | Máquina de estados | Classes |", "|---|---|---|---|---|---|---|"]
    for r in RFS:
        n = r["n"]
        st = r["estado"]["nome"]
        L.append(f"| **RF{n}** | {r['titulo']} | [png](RF{n}/RF{n}_Atividades.png) | [png](RF{n}/RF{n}_Sequencia.png) | "
                 f"[png](RF{n}/RF{n}_Componentes.png) | [{st}](RF{n}/RF{n}_MaquinaDeEstados_{st}.png) | [png](RF{n}/RF{n}_Classes.png) |")
    with open(os.path.join(DIAG, "README.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")


def planilha():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    cores = {"MySQL": "DCEBFF", "Cassandra": "E8F5E9", "Redis": "FDECEA", "OpenSearch": "FFF4E5", "S3": "F3E8FF"}
    fina = Side(style="thin", color="CBC6BE")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    cab = PatternFill("solid", fgColor="2D2438")
    wb = Workbook()

    ws = wb.active
    ws.title = "RF x Entidades x BD"
    cols = ["RF", "Requisito funcional", "Sprint", "Entidade", "Tabela / chave / índice / bucket", "Banco em produção", "Atributos", "Observação"]
    ws.append(cols)
    for r in RFS:
        for e in r["entidades"]:
            ws.append([f"RF{r['n']}", r["titulo"], r["sprint"], e[0], e[1], e[2], "\n".join(e[3]), e[4]])
    larg = [7, 48, 7, 24, 38, 14, 52, 42]
    for i, w in enumerate(larg, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = borda
    for row in ws.iter_rows(min_row=2):
        db = row[5].value
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = borda
        row[5].fill = PatternFill("solid", fgColor=cores[db])
        row[5].font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Entidades (únicas)")
    ws2.append(["Entidade", "Tabela / chave / índice / bucket", "Banco em produção", "Atributos", "RFs que usam", "Observação"])
    vistos = {}
    for r in RFS:
        for e in r["entidades"]:
            vistos.setdefault(e[0], [e, []])[1].append(f"RF{r['n']}")
    for nome, (e, rfs) in sorted(vistos.items(), key=lambda kv: (kv[1][0][2], kv[0])):
        ws2.append([nome, e[1], e[2], "\n".join(e[3]), ", ".join(rfs), e[4]])
    for i, w in enumerate([24, 40, 14, 55, 30, 42], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab
        c.border = borda
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = borda
        row[2].fill = PatternFill("solid", fgColor=cores[row[2].value])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    ws3 = wb.create_sheet("Bancos em produção")
    ws3.append(["Banco", "Papel", "Tecnologia Spring", "Padrões de uso", "Entidades", "Qtde"])
    papel = {
        "MySQL": ("Fonte da verdade transacional (ACID)", "Spring Data JPA + Hibernate 7 + Flyway; HikariCP; réplicas de leitura",
                  "Contas, perfis, esquemas, cardápio, pedidos, consentimentos LGPD, diretrizes"),
        "Cassandra": ("Séries temporais e fan-out de alta escrita", "Spring Data Cassandra (driver 4.x), TTL por tabela",
                      "Eventos de engajamento, contadores, feed, auditoria, notificações, telemetria IoT"),
        "Redis": ("Cache, sessão, rate limit, rankings", "Spring Data Redis + Lettuce; Spring Cache; Bucket4j",
                  "Allow-list de JWT, carrinho, rankings Top 10/Top 20 em ZSET, cache de viewport do globo"),
        "OpenSearch": ("Busca full-text, semântica (k-NN) e geo", "opensearch-java client + Spring AI VectorStore",
                       "Busca de esquemas/peças, visualizador mundial de restaurantes, heatmaps geohex/H3"),
        "S3": ("Objetos e mídia", "AWS SDK v2 (S3 + presigned URL) + CloudFront; SSE-KMS",
               "Fotos, artes, modelos 3D, tiles do globo, ilustrações Lottie, exportações LGPD"),
    }
    for db, (p, t, u) in papel.items():
        ents = sorted(n for n, (e, _) in vistos.items() if e[2] == db)
        ws3.append([db, p, t, u, ", ".join(ents), len(ents)])
    for i, w in enumerate([14, 36, 46, 52, 60, 8], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = borda
        row[0].fill = PatternFill("solid", fgColor=cores[row[0].value])
        row[0].font = Font(bold=True)

    ws4 = wb.create_sheet("Requisitos funcionais")
    ws4.append(["RF", "Título", "Sprint", "Tela", "Resumo", "Telas / sub-telas", "Endpoints REST", "Serviço", "Máquina de estados"])
    for r in RFS:
        ws4.append([f"RF{r['n']}", r["titulo"], r["sprint"], r["tela"], r["resumo"], "\n".join("• " + t for t in r["telas"]),
                    "\n".join(r["rotas"]), r["service"], r["estado"]["nome"]])
    for i, w in enumerate([7, 45, 7, 30, 70, 55, 50, 22, 16], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    for c in ws4[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab
    for row in ws4.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = borda
    ws4.freeze_panes = "B2"

    out = os.path.join(DOCS, "planilhas")
    os.makedirs(out, exist_ok=True)
    wb.save(os.path.join(out, "Dine_Explorer_RF_Entidades_BD.xlsx"))


def trello_json():
    dados = []
    for r in RFS:
        porbanco = {}
        for e in r["entidades"]:
            porbanco.setdefault(e[2], []).append(e[0])
        desc = (f"**{r['tela']}** · Sprint {r['sprint']}\n\n{r['resumo']}\n\n"
                f"**Backend ({SPRING})**: `{r['controller']}` → `{r['service']}`\n"
                + "\n".join(f"- `{x}`" for x in r["rotas"])
                + "\n\n**Entidades × banco**\n" + "\n".join(f"- **{db}**: {', '.join(v)}" for db, v in porbanco.items())
                + f"\n\n**Diagramas UML**: `docs/dine-explorer/diagramas/RF{r['n']}/` (atividades, sequência, componentes, "
                  f"máquina de estados {r['estado']['nome']}, classes)")
        dados.append({"n": r["n"], "nome": f"RF{r['n']}-{r['titulo']}", "sprint": r["sprint"], "desc": desc, "telas": r["telas"],
                      "erros": r["erros"]})
    with open(os.path.join(DOCS, "trello", "cards.json"), "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    os.makedirs(os.path.join(DOCS, "trello"), exist_ok=True)
    arqs = gerar_diagramas()
    indice()
    planilha()
    trello_json()
    if "--render" in sys.argv:
        render(sys.argv[sys.argv.index("--render") + 1], arqs)
    print(f"{len(arqs)} diagramas, RF1–RF{N}")
