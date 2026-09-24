"""Etapa 12 — planilhas geradas do código e do teste ponta a ponta:
  IA_por_RF_RNF.xlsx       capacidades de IA (catálogo real), onde são chamadas, por RF (Trello) e RNF
  Entidades_BD_por_RF_RNF.xlsx  entidades JPA → tabela → repositório → banco/projeções, e o que cada RF grava/lê
Uso: python3 etapa12.py results.json saida_dir"""
import collections, glob, json, os, re, sys

SCR = os.environ.get('FAI_E2E_WORKDIR', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'work'))
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
sys.path.insert(0, os.path.join(ROOT, 'scripts', 'diagramas'))
import code_model as cm  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

res = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
os.makedirs(OUT, exist_ok=True)
cat = json.load(open(os.path.join(SCR, 'aicat.json')))   # exportado por aicat.jsh (jshell com fai-application no classpath)
names = json.load(open(os.path.join(SCR, 'rf_names.json')))   # nomes dos RF (lista Requisitos Funcionais do Trello)
inv = json.load(open(os.path.dirname(__file__) + '/inventory.json'))


def rf_key(rf):
    m = re.match(r'(RN?F)(\d+)', rf)
    return (0 if m and m.group(1) == 'RF' else 1, int(m.group(2)) if m else 999, rf)


# ------------------------------------------------------------------ uso estático das capacidades
svc_dir = ROOT + '/fai-application/src/main/java/br/com/fashionai/application'
cap_files = collections.defaultdict(set)
for f in glob.glob(svc_dir + '/**/*.java', recursive=True):
    src = open(f, encoding='utf-8').read()
    for c in set(re.findall(r'AiCapability\.([A-Z_0-9]+)', src)):
        base = os.path.basename(f)[:-5]
        if base not in ('AiCapability', 'AiCatalog', 'AiEngine'):
            cap_files[c].add(base)
# controller → serviços injetados (1 ou 2 saltos)
deps = cm.DEPS


def reach(ctrl, depth=2):
    seen, frontier = set(), [ctrl]
    for _ in range(depth):
        nxt = []
        for x in frontier:
            for d in deps.get(x, []):
                if d not in seen and d.endswith('Service'):
                    seen.add(d)
                    nxt.append(d)
        frontier = nxt
    return seen


ctrl_rfs = collections.defaultdict(set)
ctrl_eps = collections.defaultdict(list)
for e in inv:
    ctrl_rfs[e['controller']].add(e['rf'])
    ctrl_eps[e['controller']].append(f"{e['verb']} {e['path']}")
cap_ctrls = collections.defaultdict(set)
for ctrl in ctrl_rfs:
    r = reach(ctrl) | {ctrl}
    for c, files in cap_files.items():
        if files & r:
            cap_ctrls[c].add(ctrl)

# ------------------------------------------------------------------ uso observado no teste E2E (INSERT em ai_inference_log)
CAPS = [c['cap'] for c in cat]
seen_rf = collections.defaultdict(lambda: collections.defaultdict(list))   # cap -> rf -> [endpoint]
prov_seen = collections.defaultdict(collections.Counter)                     # cap -> provider/result
for r in res:
    for q in r.get('sql_writes') or []:
        if 'ai_inference_log' not in q.lower():
            continue
        for c in CAPS:
            if f"'{c}'" in q:
                ep = r.get('endpoint') or r['method'] + ' ' + r['template']
                if ep not in seen_rf[c][r['rf']]:
                    seen_rf[c][r['rf']].append(ep)
                m = re.search(r"'(local|anthropic|gemini|fashn|meshy|stability|photoroom|replicate|removebg|cloudinary)'", q)
                prov_seen[c][m.group(1) if m else '?'] += 1

consent_pt = {'AI_EXTERNAL_PHOTO_PROCESSING': 'Envio de fotos a IA externa', 'AI_RECOMMENDATION': 'Recomendações por IA',
              'HISTORY_FOR_RECOMMENDATION': 'Histórico para recomendação', None: 'não exige (sem dado pessoal)'}


def prov(p):
    if not p:
        return '—'
    return f"{p['service']} · {p['model']} ({p['cost']}, ~US$ {p['costPerCall']}/chamada, {p['latency']} ms)"


HEAD = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='2F3E46')


def sheet(wb, title, head, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(head)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


# ================================================================== planilha 1: IA
wb = Workbook()
rows = []
for c in sorted(cat, key=lambda c: c['n']):
    obs = seen_rf.get(c['cap'], {})
    rows.append([c['n'], c['name'], c['cap'], c['hostRf'], c['what'], prov(c['primary']), prov(c['alternative']), prov(c['local']),
                 c['fallback'], consent_pt.get(c['consent'], c['consent']), c['quota'], c['timeout'], c['status'],
                 ', '.join(sorted(cap_files.get(c['cap'], []))),
                 ', '.join(sorted({rf for ct in cap_ctrls.get(c['cap'], []) for rf in ctrl_rfs[ct]}, key=rf_key)),
                 '; '.join(f"{rf}: {len(eps)} endpoint(s)" for rf, eps in sorted(obs.items(), key=lambda kv: rf_key(kv[0]))) or 'não acionada no teste',
                 ', '.join(f'{k}={v}' for k, v in prov_seen.get(c['cap'], {}).items()) or '—'])
sheet(wb, 'Capacidades de IA', ['Nº', 'Capacidade', 'Enum', 'RF hospedeiro (código)', 'O que faz', 'Provedor primário', 'Alternativa', 'Fallback local',
                                'Comportamento de fallback', 'Consentimento LGPD', 'Cota diária/usuário', 'Timeout (s)', 'Status', 'Serviços que chamam',
                                'RFs cujos endpoints chegam nela (estático)', 'Acionada no teste E2E (RF: endpoints)', 'Provedor usado no teste'], rows,
      (5, 26, 24, 14, 60, 40, 40, 36, 50, 24, 10, 8, 36, 36, 30, 40, 20), first=True)

# por RF (Trello)
per_rf = collections.defaultdict(lambda: collections.defaultdict(set))
for c in cat:
    for ct in cap_ctrls.get(c['cap'], []):
        for rf in ctrl_rfs[ct]:
            per_rf[rf][c['cap']].add('estático')
    for rf in seen_rf.get(c['cap'], {}):
        per_rf[rf][c['cap']].add('observado no E2E')
byname = {c['cap']: c for c in cat}
rows = []
for rf in sorted(set(per_rf) | {k for k in names}, key=rf_key):
    caps = per_rf.get(rf, {})
    if not caps:
        rows.append([rf, names.get(rf, ''), '— (sem IA)', '', '', '', '', ''])
        continue
    for cap, how in sorted(caps.items(), key=lambda kv: byname[kv[0]]['n']):
        c = byname[cap]
        rows.append([rf, names.get(rf, ''), c['name'], prov(c['primary']), prov(c['local']), consent_pt.get(c['consent'], c['consent']),
                     ' + '.join(sorted(how)), '; '.join(seen_rf.get(cap, {}).get(rf, [])[:4])])
sheet(wb, 'IA por RF', ['RF (Trello)', 'Requisito', 'Capacidade de IA', 'Provedor primário', 'Fallback local', 'Consentimento', 'Como foi mapeado',
                        'Endpoints que acionaram no teste'], rows, (10, 50, 30, 44, 40, 26, 22, 60))
rnf = [
    ['RNF5', 'Registro de toda inferência de IA', 'ai_inference_log (capacidade, provedor, modelo, latência, custo, fallback, resultado) — AiEngine grava em toda chamada, inclusive fallback local', 'MySQL ai_inference_log'],
    ['RNF6', 'Consentimento e transparência', 'Capacidades com finalidade LGPD só usam provedor externo com consentimento (PUT /api/me/consents/{purpose}); a resposta traz "explanation" (capacidade, provedor, modelo, entradas, motivo)', 'MySQL user_consents, ai_inference_log'],
    ['RNF8', 'Timeout 30 s + fallback', 'Timeout por capacidade (30 s), nova tentativa e circuit breaker; sem resposta, cai no motor local sem travar o fluxo', 'resilience4j + AiEngine'],
    ['RNF (cota)', 'Cota diária de IA por usuário', 'RateLimitPort: Redis em produção; memória local quando REDIS_ENABLED=false (este ambiente)', 'Redis (produção)'],
    ['RNF (custo)', 'Custo estimado por chamada', 'AiCatalog.tokenCost usa o uso real de tokens devolvido pelo provedor; dashboard de IA do admin (GET /api/admin/ai)', 'MySQL ai_inference_log'],
]
sheet(wb, 'IA por RNF', ['RNF', 'Requisito', 'Como o código atende', 'Onde fica'], rnf, (12, 36, 100, 30))
wb.save(os.path.join(OUT, 'IA_por_RF_RNF.xlsx'))

# ================================================================== planilha 2: entidades e bancos
repo_ent = {}
for f in glob.glob(ROOT + '/fai-domain/src/main/java/br/com/fashionai/domain/repository/*.java'):
    m = re.search(r'interface (\w+) extends JpaRepository<(\w+),', open(f).read())
    if m:
        repo_ent[m.group(2)] = m.group(1)
PROJ = {
    'wardrobe_items': 'OpenSearch índice fai-pieces (busca); S3 (fotos, estúdio, 3D)',
    'schemes': 'OpenSearch índice fai-schemes (busca); Cassandra timeline_by_user (feed); Redis contadores (views/likes); S3 (capa, card)',
    'notifications': 'Cassandra notifications_by_user (caixa de entrada)',
    'reactions': 'Redis contadores por alvo', 'comments': 'Redis contadores por alvo', 'follows': 'Redis contadores de seguidores',
    'ai_inference_log': 'Redis rate limit (cota diária de IA)', 'photos': 'S3 (arquivo original e versões)',
    'users': 'S3 (avatar, capa)', 'brand_profiles': 'S3 (logo)', 'celebrity_profiles': 'S3 (foto oficial)', 'seals': 'S3 (ícone/arte do selo)',
    'room_catalog': 'S3 (logo e arte do guarda-roupa)', 'data_export_requests': 'S3 (pacote .zip)', 'brand_logos': 'S3 (logo filtrado)',
    'backup_records': 'S3 (dump do MySQL)', 'daily_looks': 'Cassandra timeline_by_user quando publicado',
}
w_by_table = collections.defaultdict(set)
r_by_table = collections.defaultdict(set)
rf_w = collections.defaultdict(lambda: collections.defaultdict(set))
rf_r = collections.defaultdict(set)
rf_media = collections.Counter()
for r in res:
    for t, ops in r['writes'].items():
        w_by_table[t].add(r['rf'])
        rf_w[r['rf']][t].update(ops)
    for t in r['reads']:
        r_by_table[t].add(r['rf'])
        rf_r[r['rf']].add(t)
    rf_media[r['rf']] += len(r.get('media') or [])
rows = []
for ent in sorted(cm.ENTITIES):
    t = cm.TABLES.get(ent, '')
    fields = cm.ENTITIES[ent].get('fields') or []
    fnames = ', '.join((f[0] if isinstance(f, (list, tuple)) else str(f)) for f in fields[:14])
    rows.append([ent, t, repo_ent.get(ent, '—'), 'MySQL (fonte da verdade)', PROJ.get(t, '—'),
                 ', '.join(sorted(w_by_table.get(t, []), key=rf_key)) or '—', ', '.join(sorted(r_by_table.get(t, []), key=rf_key)) or '—', fnames])
wb2 = Workbook()
sheet(wb2, 'Entidades', ['Entidade JPA', 'Tabela MySQL', 'Repositório', 'Banco', 'Projeções / outros bancos (produção)', 'RFs que gravam (E2E)',
                         'RFs que leem (E2E)', 'Campos (primeiros)'], rows, (26, 28, 32, 22, 50, 40, 40, 80), first=True)
rows = []
for rf in sorted(set(rf_w) | set(rf_r) | set(names), key=rf_key):
    wt = rf_w.get(rf, {})
    rows.append([rf, names.get(rf, ''), '; '.join(f"{t} ({'/'.join(sorted(o))})" for t, o in sorted(wt.items())) or '—',
                 ', '.join(sorted(rf_r.get(rf, []))) or '—', rf_media.get(rf, 0),
                 '; '.join(sorted({PROJ[t] for t in wt if t in PROJ})) or '—'])
sheet(wb2, 'BD por RF', ['RF (Trello) / RNF', 'Requisito', 'Tabelas gravadas no teste (operação)', 'Tabelas lidas no teste', 'Arquivos de mídia gravados',
                         'Projeções em produção'], rows, (14, 50, 80, 70, 12, 60))
stores = [
    ['MySQL 8', 'Fonte da verdade de todas as entidades (Flyway V1–V21)', 'sempre ligado', 'fai-infrastructure/persistence-mysql'],
    ['Redis', 'Contadores (views, likes), rate limit e cota de IA', 'REDIS_ENABLED (desligado aqui → memória local)', 'fai-infrastructure/cache-redis'],
    ['Cassandra', 'timeline_by_user (feed) e notifications_by_user', 'CASSANDRA_ENABLED (desligado aqui → consulta no MySQL)', 'fai-infrastructure/persistence-cassandra'],
    ['OpenSearch', 'Índices fai-pieces e fai-schemes (busca)', 'OPENSEARCH_ENABLED (desligado aqui → busca no MySQL)', 'fai-infrastructure/search-opensearch'],
    ['S3 / MinIO', 'Mídia: fotos, cards, logos, modelos 3D, exportações, backups', 'STORAGE_TYPE=s3 (aqui: disco local)', 'fai-infrastructure/storage-s3'],
]
sheet(wb2, 'Bancos', ['Banco', 'Guarda', 'Neste ambiente', 'Módulo'], stores, (14, 60, 50, 40))
wb2.save(os.path.join(OUT, 'Entidades_BD_por_RF_RNF.xlsx'))

# ================================================================== resumo em Markdown
L = ['# Etapa 12 — IA e entidades/bancos por RF e RNF', '',
     'Gerado de `AiCatalog`/`AiCapability` (código), do grafo de injeção controller → serviço, das entidades JPA e do teste ponta a ponta '
     '(`general_log` do MySQL). Numeração dos RF: **Trello**. Planilhas: `IA_por_RF_RNF.xlsx` e `Entidades_BD_por_RF_RNF.xlsx`.', '',
     '## Capacidades de IA', '', '| Nº | Capacidade | RF hospedeiro | Primário | Fallback local | Consentimento | Acionada no E2E |', '|---|---|---|---|---|---|---|']
for c in sorted(cat, key=lambda c: c['n']):
    obs = seen_rf.get(c['cap'], {})
    L.append(f"| {c['n']} | {c['name']} | {c['hostRf']} | {c['primary']['service'] + ' · ' + c['primary']['model'] if c['primary'] else '—'} | "
             f"{c['local']['service'] if c['local'] else '—'} | {consent_pt.get(c['consent'], c['consent'])} | {', '.join(sorted(obs, key=rf_key)) or '—'} |")
L += ['', '> No teste deste ambiente não há chave de IA configurada: toda chamada caiu no motor local e ficou registrada em `ai_inference_log` '
      '(provedor `local`, `fallback=true`).', '', '## IA por RF (Trello)', '', '| RF | Requisito | Capacidades de IA |', '|---|---|---|']
for rf in sorted(set(per_rf) | set(names), key=rf_key):
    caps = per_rf.get(rf, {})
    L.append(f"| {rf} | {names.get(rf, '')} | {', '.join(byname[c]['name'] for c in sorted(caps, key=lambda c: byname[c]['n'])) or '—'} |")
L += ['', '## Bancos por RF (o que o teste gravou)', '', '| RF | Tabelas gravadas | Mídia | Projeções em produção |', '|---|---|---|---|']
for rf in sorted(set(rf_w) | set(names), key=rf_key):
    wt = rf_w.get(rf, {})
    L.append(f"| {rf} | {', '.join(sorted(wt)) or '—'} | {rf_media.get(rf, 0)} | {'; '.join(sorted({PROJ[t] for t in wt if t in PROJ})) or '—'} |")
L += ['', '## Bancos', '', '| Banco | Guarda | Neste ambiente |', '|---|---|---|'] + [f'| {a} | {b} | {c} |' for a, b, c, _ in stores]
open(os.path.join(OUT, 'ETAPA12_IA_E_ENTIDADES_POR_RF.md'), 'w').write('\n'.join(L) + '\n')
print('capacidades', len(cat), 'entidades', len(cm.ENTITIES), 'rfs', len(set(rf_w) | set(names)))
