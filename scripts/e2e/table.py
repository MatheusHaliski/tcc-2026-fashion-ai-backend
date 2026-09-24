"""Tabela do teste ponta a ponta por RF: RF, CA, endpoint, status, banco em que salvou, e se o frontend exibe (GET que lê do banco).
Uso: python3 table.py results.json saida_dir"""
import json, re, sys, os, collections

res = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
os.makedirs(OUT, exist_ok=True)
inv = json.load(open(os.path.join(os.path.dirname(__file__), 'inventory.json')))
fmap = json.load(open(os.path.join(os.path.dirname(__file__), 'frontend_map.json')))


def rx(tpl):
    return re.compile('^' + re.sub(r'\\\{[^}]+\\\}', r'[^/]+', re.escape(tpl)) + '$')


pats = sorted([(e, rx(e['path'])) for e in inv], key=lambda t: (t[0]['path'].count('{'), -len(t[0]['path'])))


def endpoint(method, path):
    p = path.split('?')[0]
    for e, r in pats:
        if e['verb'] == method and r.match(p):
            return e['verb'] + ' ' + e['path']
    return method + ' ' + p


for r in res:
    r['endpoint'] = endpoint(r['method'], r['path'])

# telas que chamam cada endpoint
def pages(ep):
    fs = fmap.get(ep) or []
    return [f.replace('app/(app)/', '').replace('app/(auth)/', '').replace('/page.tsx', '').replace('components/', 'componente ') or '/' for f in fs]


# para cada tabela: GETs (com tela) que a leem, com o nº de tabelas lidas (quanto menos, mais específico)
readers = collections.defaultdict(dict)
for r in res:
    if r['method'] == 'GET' and r['ok'] and pages(r['endpoint']):
        for t in r['reads']:
            n = readers[t].get(r['endpoint'])
            readers[t][r['endpoint']] = min(n, len(r['reads'])) if n else len(r['reads'])

PREF = {'wardrobe_items': ['GET /api/pieces/{id}', 'GET /api/me/closet'], 'schemes': ['GET /api/schemes/{id}'],
        'dna_schemes': ['GET /api/me/dna-schemes'], 'room_layouts': ['GET /api/me/room'], 'challenge_instances': ['GET /api/challenges/{id}']}
SIDE = ('audit_log', 'ai_inference_log', 'notifications', 'notifications_outbox', 'item_embeddings', 'refresh_tokens')


def segs(ep):
    return [x for x in ep.split(' ', 1)[-1].split('/') if x and not x.startswith('{')]


def best_readers(write_ep, tables):
    """GETs com tela que releem as tabelas gravadas: primeiro os da mesma rota (recurso), depois os mais específicos."""
    ws = segs(write_ep)
    cands = {}
    for i, t in enumerate(sorted(tables, key=lambda t: (t in SIDE, t))):
        for g, n in readers.get(t, {}).items():
            common = 0
            for a, b in zip(ws, segs(g)):
                if a != b:
                    break
                common += 1
            pref = PREF.get(t, [])
            same = g.split(' ', 1)[-1].rsplit('/', 1)[0] == write_ep.split(' ', 1)[-1] or g.split(' ', 1)[-1] == write_ep.split(' ', 1)[-1]
            score = (t in SIDE, pref.index(g) if g in pref else 9, 0 if same else 1, -common, n)
            if g not in cands or score < cands[g]:
                cands[g] = score
    return [g for g, _ in sorted(cands.items(), key=lambda kv: kv[1])]


MEDIA = 'storage de mídia (local; S3 em produção)'


def banco(r):
    if r['status'] >= 400 and not r['writes'] and not r['media']:
        return '— (nada gravado: requisição recusada)'
    parts = []
    if r['writes']:
        parts.append('MySQL: ' + ', '.join(f"{t} ({'/'.join(ops)})" for t, ops in sorted(r['writes'].items())))
    if r['media']:
        parts.append(MEDIA + f": {len(r['media'])} arquivo(s)")
    if not parts:
        return '— (não grava)' if r['method'] == 'GET' else '— (não gravou)'
    return '; '.join(parts)


def exibe(r):
    pg = pages(r['endpoint'])
    if r['status'] >= 400:
        code = r.get('error')
        try:
            code = code or json.loads(r.get('resp') or '{}').get('code')
        except (ValueError, AttributeError):
            pass
        return f"não se aplica — recusado com {r['status']} {code or ''} (regra de negócio testada)".replace('  ', ' ') + ('' if r['ok'] else ' · INESPERADO')
    if r['method'] == 'GET':
        src = ('lê MySQL: ' + ', '.join(r['reads'][:6])) if r['reads'] else 'sem leitura no MySQL (cálculo/cache/arquivo)'
        return ('sim — ' + '; '.join(pg[:3]) + ' · ' + src) if pg else ('API sem tela própria · ' + src)
    tabs = [t for t in r['writes'] if t not in ('audit_log', 'ai_inference_log', 'notifications_outbox')]
    gets = best_readers(r['endpoint'], tabs)
    via = ('; tela: ' + '; '.join(pg[:2])) if pg else ''
    if gets:
        return f"sim — relido por {', '.join(gets[:2])}{via}"
    if r['media']:
        return 'sim — arquivo servido em /media' + via
    return ('ação sem releitura' + via) if pg else 'ação de API (sem tela)'


def status_txt(r):
    s = r['status']
    return f"{s}" + ('' if r['ok'] else ' ✗')


rows = []
for r in res:
    rows.append(dict(rf=r['rf'], ca=r['ca'], endpoint=r['endpoint'], passo=r['desc'], usuario=r['who'], status=r['status'], ok=r['ok'],
                     banco=banco(r), exibe=exibe(r), ms=r['ms'], erro=r.get('error')))
json.dump(rows, open(os.path.join(OUT, 'tabela.json'), 'w'), ensure_ascii=False, indent=1)

# ordem: RF numérico, depois RNF
def rf_key(rf):
    m = re.match(r'(RN?F)(\d+)', rf)
    return (0 if m and m.group(1) == 'RF' else 1, int(m.group(2)) if m else 999, rf)


by_rf = collections.defaultdict(list)
for x in rows:
    by_rf[x['rf']].append(x)
eps = {x['endpoint'] for x in rows}
ok = sum(x['ok'] for x in rows)
L = ['# Teste ponta a ponta dos endpoints — tabela por RF', '',
     f'Execução real contra a API (Spring Boot) e o MySQL local. **{len(rows)} passos, {ok} com o status esperado, {len(eps)} endpoints distintos** '
     f'(de {len(inv)} no inventário do backend).', '',
     'Como cada coluna foi obtida:', '',
     '- **Status**: código HTTP devolvido. Alguns passos testam regras de negócio e esperam 4xx (ex.: 409 limite, 403 sem permissão); o status real aparece na tabela.',
     '- **Salvou em qual banco?**: comandos `INSERT/UPDATE/DELETE` que a própria aplicação executou durante a chamada, lidos do `general_log` do MySQL (filtrado pelo usuário `fashionai`), mais arquivos novos no storage de mídia.',
     '- **Exibiu no frontend?**: para GET, as telas do Next.js que chamam o endpoint (varredura do código) e as tabelas que o GET leu; para escritas, qual GET com tela relê a tabela gravada.',
     '- Neste ambiente Redis, Cassandra e OpenSearch estão desligados (`*_ENABLED=false`): cache, timeline e busca caem no MySQL; o storage de mídia é o disco local (S3 em produção).', '']
for rf in sorted(by_rf, key=rf_key):
    xs = by_rf[rf]
    L.append(f'## {rf}  ({len(xs)} passos, {len({x["endpoint"] for x in xs})} endpoints)')
    L.append('')
    L.append('| CA | Endpoint | Passo | Usuário | Status | Salvou em qual banco? | Exibiu no frontend (GET buscou do banco)? |')
    L.append('|----|----------|-------|---------|--------|-----------------------|-------------------------------------------|')
    for x in xs:
        cell = lambda v: str(v).replace('|', '\\|').replace('\n', ' ')
        L.append(f"| {cell(x['ca'])} | `{cell(x['endpoint'])}` | {cell(x['passo'])} | {cell(x['usuario'])} | {x['status']}{'' if x['ok'] else ' ✗'} | {cell(x['banco'])} | {cell(x['exibe'])} |")
    L.append('')
open(os.path.join(OUT, 'TABELA_ENDPOINTS_POR_RF.md'), 'w').write('\n'.join(L) + '\n')

# planilha
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = 'Endpoints por RF'
    head = ['RF', 'CA', 'Endpoint', 'Passo', 'Usuário', 'Status HTTP', 'Esperado?', 'Salvou em qual banco?', 'Exibiu no frontend (GET buscou do banco)?', 'ms']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2F3E46')
    for x in sorted(rows, key=lambda x: rf_key(x['rf'])):
        ws.append([x['rf'], x['ca'], x['endpoint'], x['passo'], x['usuario'], x['status'], 'sim' if x['ok'] else 'NÃO', x['banco'], x['exibe'], x['ms']])
    for col, w in zip('ABCDEFGHIJ', (8, 8, 52, 46, 18, 11, 10, 60, 70, 8)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    s2 = wb.create_sheet('Resumo por RF')
    s2.append(['RF', 'Passos', 'Endpoints', 'Com status esperado', 'Tabelas MySQL gravadas'])
    for c in s2[1]:
        c.font = Font(bold=True)
    for rf in sorted(by_rf, key=rf_key):
        xs = by_rf[rf]
        tabs = sorted({t for r in res if r['rf'] == rf for t in r['writes']})
        s2.append([rf, len(xs), len({x['endpoint'] for x in xs}), sum(x['ok'] for x in xs), ', '.join(tabs)])
    s2.column_dimensions['E'].width = 120
    wb.save(os.path.join(OUT, 'TABELA_ENDPOINTS_POR_RF.xlsx'))
except ImportError:
    print('openpyxl ausente: só o .md foi gerado')
print('passos', len(rows), 'ok', ok, 'endpoints', len(eps))
