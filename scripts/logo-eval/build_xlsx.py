"""
Planilha de testes do RF4 (tipo, cor, logo e marca) a partir de dataset.json e do resultado de run_eval.py.

Uso: python3 build_xlsx.py resultados.json pasta-das-imagens saida.xlsx
"""
import io
import json
import sys
from collections import Counter

import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

from pipeline import decide

RES, IMG_DIR, OUT = sys.argv[1:4]
data = json.load(open("dataset.json", encoding="utf-8"))
R = json.load(open(RES, encoding="utf-8"))

# O que conta como acerto de tipo e de cor (a taxonomia do FashionAI é mais fina que o rótulo manual).
TIPO = {
    "tênis": ("shoes_piece", {"casual_sneakers", "running_shoes", "training_shoes", "basketball_shoes", "skate_shoes", "high_top_sneakers"}),
    "camiseta": ("upper_piece", {"t_shirt"}), "camisa polo": ("upper_piece", {"polo_shirt"}), "camisa": ("upper_piece", {"shirt"}),
    "camisa esportiva": ("upper_piece", {"t_shirt", "polo_shirt", "shirt", "tank_top"}), "regata": ("upper_piece", {"tank_top"}),
    "moletom": ("upper_piece", {"sweatshirt", "hoodie"}), "jaqueta": ("upper_piece", {"jacket", "windbreaker", "coat", "parka"}),
    "colete": ("upper_piece", {"vest"}), "boné": ("accessory_piece", {"cap"}), "chapéu": ("accessory_piece", {"hat"}),
    "bolsa": ("accessory_piece", {"handbag", "tote_bag", "crossbody_bag", "backpack"}), "jeans": ("lower_piece", {"jeans"}),
}
COR = {
    "preto": {"black", "charcoal", "washed_black", "dark_gray"}, "branco": {"white", "off_white", "ivory", "cream"},
    "cinza": {"light_gray", "gray", "dark_gray", "silver", "charcoal"},
    "azul": {"blue", "light_blue", "sky_blue", "cobalt", "denim", "navy", "teal"}, "azul-marinho": {"navy", "blue", "cobalt"},
    "vermelho": {"red", "crimson", "burgundy", "maroon"}, "verde": {"green", "olive", "military_green", "forest_green", "mint", "sage", "emerald"},
    "verde-limão": {"green", "yellow", "mint"}, "amarelo": {"yellow", "mustard", "gold", "butter"},
    "roxo": {"purple", "violet", "lilac", "lavender", "plum"}, "marrom": {"brown", "chocolate", "camel", "tan"},
    "cáqui": {"beige", "tan", "camel", "taupe", "olive"}, "bege": {"beige", "cream", "tan", "camel", "off_white", "ivory"},
}
ESTADO = {"CONFIRMADA": "Marca confirmada", "POSSIVEL": "Possível marca — confirmar", "LOGO_SEM_MARCA": "Logo sem marca conhecida — pedir a marca",
          "SEM_LOGO": "Sem logo"}

items = [i for i in data["itens"] if not i.get("excluded")]
rows, cnt = [], Counter()
for it in items:
    r = R[str(it["n"])]; cur = r["rf4"]
    res = decide([(t, c) for t, c in r["textosConf"]], np.array(r["probs"]))
    cat, subs = TIPO[it["tipo"]]
    cat_ok = cur.get("category") == cat; sub_ok = cur.get("subcategory") in subs
    cor_ok = cur.get("color") in COR[it["cor"]]
    exp = it["marca"]; k = it["classe"]
    cur_brand = cur.get("brand") if cur.get("brand") not in (None, "", "Sem marca") else None
    cur_ok = (cur_brand == exp) if k == "A" else (cur_brand is None)
    if k == "A":
        if res.marca == exp:
            v = "Acerto (confirmada)" if res.estado == "CONFIRMADA" else "Acerto (pede confirmação)"
        elif res.marca:
            v = "Erro: marca errada"
        else:
            v = "Não nomeou: pede a marca" if res.estado == "LOGO_SEM_MARCA" else "Não achou o logo"
    else:
        v = "Correto (não inventou marca)" if not res.marca else ("Falso positivo confirmado" if res.estado == "CONFIRMADA" else "Falso positivo (pede confirmação)")
    cnt.update({f"cat_{cat_ok}", f"sub_{sub_ok}", f"cor_{cor_ok}", f"{k}:{v}", f"rf4marca_{k}_{cur_ok}"})
    filled = all(cur.get(f) not in (None, "") for f in ("category", "subcategory", "color", "brand"))
    cnt[f"preenche_{filled}"] += 1
    rows.append((it, cur, res, cat_ok, sub_ok, cor_ok, cur_brand, v, r))

wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F2A44"); OK = PatternFill("solid", fgColor="DCF2E3")
BAD = PatternFill("solid", fgColor="F9DEDC"); MID = PatternFill("solid", fgColor="FFF1CC"); WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=ws.max_row, column=i); c.font = H; c.fill = HF; c.alignment = WRAP
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=3)


# ---------- Resumo ----------
ws = wb.active; ws.title = "Resumo"
nA = sum(1 for i in items if i["classe"] == "A"); nBC = len(items) - nA
acertoA = cnt["A:Acerto (confirmada)"] + cnt["A:Acerto (pede confirmação)"]
fpBC = sum(v for kk, v in cnt.items() if kk.startswith(("B:", "C:", "D:")) and "Falso" in kk)
lines = [
    ("Testes do RF4 — tipo, cor, logo e marca (27/09/2026)", ""),
    ("Fotos testadas", len(items)),
    ("Fotos com marca de moda visível (classe A)", nA),
    ("Fotos sem marca de moda (classes B, C e D)", nBC),
    ("", ""),
    ("RF4 atual, backend local (IA remota desligada neste ambiente)", ""),
    ("Categoria certa", f"{cnt['cat_True']} de {len(items)}"),
    ("Subcategoria certa", f"{cnt['sub_True']} de {len(items)}"),
    ("Cor certa", f"{cnt['cor_True']} de {len(items)}"),
    ("Marca certa nas fotos com marca", f"{cnt['rf4marca_A_True']} de {nA}"),
    ("Todos os campos preenchidos", f"{cnt['preenche_True']} de {len(items)}"),
    ("", ""),
    ("Pipeline proposto (OCR + símbolo em recortes ampliados, local)", ""),
    ("Marca certa e confirmada pelo texto da peça", f"{cnt['A:Acerto (confirmada)']} de {nA}"),
    ("Marca certa sugerida, a pessoa confirma", f"{cnt['A:Acerto (pede confirmação)']} de {nA}"),
    ("Total de marca certa", f"{acertoA} de {nA}"),
    ("Marca errada", f"{cnt['A:Erro: marca errada']} de {nA}"),
    ("Logo visto, marca não nomeada: o app pede a marca", f"{cnt['A:Não nomeou: pede a marca']} de {nA}"),
    ("Logo não achado", f"{cnt['A:Não achou o logo']} de {nA}"),
    ("Sugeriu marca em foto sem marca de moda", f"{fpBC} de {nBC}"),
    ("", ""),
    ("Leitura", "O RF4 atual, sem IA remota, não nomeia marca nenhuma. O pipeline proposto nomeia a marca quando ela está escrita "
                "(etiqueta, texto no tênis) e sugere quando é só símbolo; nunca termina em silêncio. Símbolo pequeno em polo e camisa "
                "esportiva ainda falha localmente: é o caso que pede o reconhecedor remoto (IA de visão com os recortes ou Google Cloud Vision)."),
]
for a, b in lines:
    ws.append([a, b])
ws.column_dimensions["A"].width = 58; ws.column_dimensions["B"].width = 90
ws["A1"].font = Font(bold=True, size=14)
for rr in (6, 13):
    ws.cell(row=rr, column=1).font = Font(bold=True)
ws.cell(row=len(lines), column=2).alignment = WRAP

# ---------- Testes ----------
ws = wb.create_sheet("Testes")
cols = ["#", "Foto", "URL da imagem", "Fonte original (Flickr)", "Autor", "Licença", "Tipo de peça (esperado)", "Tipo detectado (RF4 atual)", "Tipo certo?",
        "Cor (esperada)", "Cor detectada (RF4 atual)", "Cor certa?", "Classe", "Marca esperada", "Marca detectada (RF4 atual)",
        "Estado (pipeline proposto)", "Marca detectada (pipeline proposto)", "Como achou", "Confiança", "Resultado da marca",
        "Texto lido na peça (OCR)", "Símbolos mais prováveis (CLIP)", "Tempo RF4 (ms)", "Tempo pipeline (ms)"]
header(ws, cols, [5, 16, 34, 34, 18, 16, 16, 26, 9, 12, 16, 9, 7, 18, 18, 30, 20, 16, 10, 26, 36, 36, 10, 10])
for idx, (it, cur, res, cat_ok, sub_ok, cor_ok, cur_brand, v, r) in enumerate(rows, start=2):
    ws.append([it["n"], "", it["imageUrl"], it["pageUrl"], it["author"], "CC BY 2.0", it["tipo"],
               f"{cur.get('category')} / {cur.get('subcategory')}", "Sim" if (cat_ok and sub_ok) else ("Só a categoria" if cat_ok else "Não"),
               it["cor"], cur.get("color"), "Sim" if cor_ok else "Não", it["classe"], it["marca"] or "— (sem marca de moda)",
               cur_brand or "Sem marca", ESTADO[res.estado], res.marca or "—", res.fonte or "—", res.confianca or "",
               v, " · ".join(t for t, _ in r["textosConf"][:8]), " · ".join(f"{b} {p:.2f}" for b, p in res.clip),
               r["rf4Ms"], r["pipelineMs"]])
    ws.cell(row=idx, column=3).hyperlink = it["imageUrl"]; ws.cell(row=idx, column=4).hyperlink = it["pageUrl"]
    for c in range(1, len(cols) + 1):
        ws.cell(row=idx, column=c).alignment = WRAP
    for col, good in ((9, cat_ok and sub_ok), (12, cor_ok)):
        ws.cell(row=idx, column=col).fill = OK if good else BAD
    ws.cell(row=idx, column=20).fill = OK if v.startswith(("Acerto", "Correto")) else (MID if v.startswith(("Não nomeou", "Falso positivo (pede")) else BAD)
    im = Image.open(f"{IMG_DIR}/{it['imageId']}.jpg").convert("RGB"); im.thumbnail((110, 110))
    buf = io.BytesIO(); im.save(buf, "JPEG", quality=80); buf.seek(0)
    xi = XlImage(buf); ws.add_image(xi, f"B{idx}"); ws.row_dimensions[idx].height = 86
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

# ---------- Calibração ----------
ws = wb.create_sheet("Calibração")
header(ws, ["Limiar do símbolo", "Folga sobre a 2ª marca", "Marca certa (A)", "Marca errada (A)", "Sugestão falsa (B/C/D)"], [18, 20, 16, 16, 22])
for pos, mar in ((0.35, 0.15), (0.5, 0.2), (0.6, 0.25), (0.7, 0.3), (0.8, 0.3)):
    c = Counter()
    for it in items:
        r = R[str(it["n"])]; res = decide([(t, cc) for t, cc in r["textosConf"]], np.array(r["probs"]), possible=pos, margin=mar)
        if it["classe"] == "A":
            c["ok" if res.marca == it["marca"] else ("err" if res.marca else "none")] += 1
        elif res.marca:
            c["fp"] += 1
    ws.append([pos, mar, f"{c['ok']} de {nA}", f"{c['err']} de {nA}", f"{c['fp']} de {nBC}"])
ws.append([]); ws.append(["Ponto escolhido: 0,60 e 0,25. Foi escolhido nestas mesmas fotos, então o número é otimista: precisa ser confirmado num conjunto novo."])

# ---------- Excluídas ----------
ws = wb.create_sheet("Excluídas")
header(ws, ["#", "URL da imagem", "Motivo"], [5, 60, 50])
for it in data["itens"]:
    if it.get("excluded"):
        ws.append([it["n"], it["imageUrl"], it["excluded"]])

# ---------- Método ----------
ws = wb.create_sheet("Método")
for line in [
    "Fonte das imagens: Open Images V7 (Google). Cada foto é do Flickr, sob licença CC BY 2.0; o autor e a página original estão na aba Testes.",
    "Seleção: todas as 100 fotos das partições de validação e teste com rótulo humano de peça (camiseta, polo, tênis, boné, jaqueta, jeans…) e de logo ou marca.",
    "Rótulos esperados (tipo, cor, marca, classe): feitos à mão, olhando cada foto. 11 fotos foram excluídas (aba Excluídas).",
    "Classe A: marca de moda visível. B: logo de outra entidade (time, empresa, evento). C: sem marca. D: logo que o rotulador não reconheceu.",
    "RF4 atual: POST /api/pieces/analysis num backend local isolado, com a IA remota desligada. Em produção a análise usa IA de visão; "
    "esse caminho não foi executado aqui para não usar as chaves de produção sem autorização.",
    "Pipeline proposto: scripts/logo-eval/pipeline.py. OCR RapidOCR (Apache-2.0) e CLIP ViT-B/32 em ONNX (MIT), rodando na máquina, em 14 recortes por foto.",
    "Regra: texto lido confirma a marca; símbolo só sugere ('Possível marca: X — confirmar') e precisa vencer com folga.",
    "Reproduzir: API=… TOKEN=… CLIP_DIR=… CACHE=… python3 scripts/logo-eval/run_eval.py resultados.json; depois build_xlsx.py.",
]:
    ws.append([line])
ws.column_dimensions["A"].width = 160
wb.save(OUT)
print("ok", OUT, dict(cnt))
