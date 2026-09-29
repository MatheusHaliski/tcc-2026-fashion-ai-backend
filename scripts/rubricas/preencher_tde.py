"""Preenche o TDE a partir da seção 4 usando o Markdown versionável.

O corte é localizado pelo título, nunca por um índice fixo. Os elementos OOXML
anteriores ao título são mantidos no mesmo documento; a validação compara sua
representação canônica antes e depois da gravação.
"""
from __future__ import annotations

import hashlib
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "TDEword" / "DOCUMENTACAO_FASHION_AI_VERSAO_FINAL.docx"
BACKUP = ROOT / "backups" / "DOCUMENTACAO_FASHION_AI_VERSAO_FINAL.original.docx"
AUDIT = ROOT / "TDEword" / "AUDITORIA_ORIGINAL_ANTES_SECAO_4.txt"
SOURCE = ROOT / "docs" / "rubricas" / "TDE_SECAO_4_RF25_RF39.md"
SECTION_4 = "RELAÇÃO DE ATORES / USUÁRIOS"


def find_boundary(document: Document) -> int:
    body = document._element.body
    for index, element in enumerate(body.iterchildren()):
        if etree.QName(element).localname != "p":
            continue
        text = "".join(node.text or "" for node in element.iter(qn("w:t")))
        normalized = text.upper()
        if normalized.strip() == SECTION_4 or normalized.startswith("4. RELAÇÃO DE ATORES"):
            return index
    raise RuntimeError(f"Título da seção 4 não encontrado: {SECTION_4}")


def canonical_prefix(document: Document, boundary: int) -> str:
    body = document._element.body
    payload = b"".join(
        etree.tostring(element, method="c14n", with_comments=True)
        for element in list(body.iterchildren())[:boundary]
    )
    return hashlib.sha256(payload).hexdigest()


def extract_audit(document: Document) -> str:
    lines = [
        "AUDITORIA TEXTUAL DO DOCX ORIGINAL",
        "Arquivo: TDEword/DOCUMENTACAO_FASHION_AI_VERSAO_FINAL.docx",
        "Observação: extração anterior à edição; tabelas aparecem por célula.",
        "",
    ]
    body = document._element.body
    for index, element in enumerate(body.iterchildren()):
        tag = etree.QName(element).localname
        texts = [(node.text or "").strip() for node in element.iter(qn("w:t")) if (node.text or "").strip()]
        if texts:
            lines.append(f"[{index:03d} {tag}] " + " | ".join(texts))
    return "\n".join(lines) + "\n"


def style_by_id(document: Document, style_id: str):
    return next((style for style in document.styles if style.style_id == style_id), None)


def clean_markdown(value: str) -> str:
    value = re.sub(r"`([^`]+)`", r"\1", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"\1", value)
    return value.replace("—", "—").strip()


def add_paragraph(document: Document, text: str, *, style=None, bold_prefix: bool = False) -> None:
    paragraph = document.add_paragraph(style=style)
    if bold_prefix and ":" in text:
        prefix, rest = text.split(":", 1)
        paragraph.add_run(prefix + ":").bold = True
        paragraph.add_run(rest)
    else:
        paragraph.add_run(text)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=columns)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for col, value in enumerate(rows[0]):
        run = table.rows[0].cells[col].paragraphs[0].add_run(clean_markdown(value))
        run.bold = True
    for row in rows[2:] if len(rows) > 1 and set("".join(rows[1])) <= set("-: ") else rows[1:]:
        cells = table.add_row().cells
        for col in range(columns):
            cells[col].text = clean_markdown(row[col] if col < len(row) else "")


def append_markdown(document: Document, markdown: str) -> None:
    title_1 = style_by_id(document, "Ttulo1")
    title_2 = style_by_id(document, "Ttulo2") or title_1
    normal = style_by_id(document, "Normal")
    lines = markdown.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 4."))
    index = start
    while index < len(lines):
        raw = lines[index].rstrip()
        if not raw:
            index += 1
            continue
        if raw.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue
        if raw.startswith("## "):
            add_paragraph(document, clean_markdown(raw[3:]).upper(), style=title_1)
        elif raw.startswith("### "):
            add_paragraph(document, clean_markdown(raw[4:]), style=title_2)
        elif re.match(r"^\d+\. \*\*", raw):
            add_paragraph(document, clean_markdown(raw), style=normal, bold_prefix=True)
        elif re.match(r"^\d+\. ", raw):
            add_paragraph(document, clean_markdown(raw), style=normal)
        elif raw.startswith("- "):
            add_paragraph(document, "• " + clean_markdown(raw[2:]), style=normal, bold_prefix=True)
        elif raw.startswith("> "):
            add_paragraph(document, clean_markdown(raw[2:]), style=normal)
        else:
            add_paragraph(document, clean_markdown(raw), style=normal)
        index += 1


def main() -> None:
    if not DOCX.is_file():
        raise FileNotFoundError(DOCX)
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    # O backup representa exatamente a entrada desta execução. Reutilizar um
    # backup antigo faria uma segunda execução descartar revisões posteriores.
    shutil.copy2(DOCX, BACKUP)

    original = Document(BACKUP)
    # A auditoria versionada é a fotografia do documento recebido originalmente
    # e não deve ser reescrita quando o gerador for executado novamente.
    if not AUDIT.exists():
        AUDIT.write_text(extract_audit(original), encoding="utf-8")
    boundary = find_boundary(original)
    before_hash = canonical_prefix(original, boundary)

    document = Document(BACKUP)
    body = document._element.body
    boundary = find_boundary(document)
    for element in list(body.iterchildren())[boundary:]:
        if etree.QName(element).localname != "sectPr":
            body.remove(element)
    append_markdown(document, SOURCE.read_text(encoding="utf-8"))
    document.save(DOCX)

    reopened = Document(DOCX)
    new_boundary = find_boundary(reopened)
    after_hash = canonical_prefix(reopened, new_boundary)
    if before_hash != after_hash:
        raise RuntimeError("A estrutura OOXML anterior à seção 4 foi alterada")
    print(f"Seção 4 localizada no elemento {boundary}; prefixo preservado: {before_hash}")
    print(f"DOCX salvo e reaberto: {DOCX}")


if __name__ == "__main__":
    main()
